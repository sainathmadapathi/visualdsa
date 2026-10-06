"""Your own practice sheet: read a list of problems from a file, a link or recognised text,
match each row to a built-in lab, and validate the labs a learner builds for the rest.

Nothing here invents problem content. A row carries only what the sheet says (a title, a link,
a difficulty, a topic); statements, parameters and expected outputs come from the learner."""
import ast
import contextvars
import csv
import functools
import html.parser
import http.client
import io
import ipaddress
import json
import keyword
import re
import socket
import ssl
import time
import urllib.error
import urllib.parse
import urllib.request
import urllib.robotparser

import leetcode
import zipfile

MAX_ROWS = 1000  # Large public sheets (Striver's A2Z has about 450 problems) must fit.
MAX_CASES = 8
MAX_FETCH = 3_000_000
FETCH_SECONDS = 30  # A whole fetch (connect, headers, body), however slowly the site drips it.
MAX_XLSX = 4_000_000  # Unpacked bytes of an Excel file: a 1,000-row sheet is under 1 MB; a zip bomb is far more.
AGENT = "VisualDSA-sheet-import/1.0"
JUDGES = ("leetcode.com", "geeksforgeeks.org", "naukri.com", "codingninjas.com", "hackerrank.com", "codeforces.com",
          "interviewbit.com", "spoj.com", "codechef.com", "takeuforward.org", "neetcode.io", "atcoder.jp", "cses.fi")
# Apostrophes may sit inside a link (takeUforward's "kadane's-algorithm"); quotes and punctuation never end one.
URL = re.compile(r"https?://[^\s<>\"\])]*[^\s<>\"\])'.,;:!?]", re.I)

# Built-in labs and the sheet titles that name the same contract ("same") or a close relative whose
# details differ ("close": e.g. LeetCode's in-place or 1-indexed version). Matching is exact on
# normalised titles and link slugs; anything else stays unmatched for the learner to decide.
ALIASES = {
    "two-sum": (["two sum", "2 sum", "pair with given sum", "two sum problem"], []),
    "contains-duplicate": (["contains duplicate", "check for duplicates"], []),
    "pair-count": (["count matching pairs", "count pairs with given sum", "number of pairs with given sum"], []),
    "frequency-map": (["count frequencies", "frequency of array elements", "count frequency of each element", "frequencies of array elements"], []),
    "unique-values": (["keep unique values", "distinct elements"], []),
    "intersection": (["array intersection", "intersection of two arrays"], ["intersection of two arrays ii"]),
    "first-unique": (["first unique character", "first unique character in a string", "first non repeating character"], []),
    "valid-anagram": (["valid anagram", "check anagram", "anagram"], []),
    "palindrome": (["is it a palindrome", "palindrome string", "check palindrome"], ["valid palindrome", "palindrome"]),
    "reverse-string": (["reverse a string", "reverse string"], []),
    "move-zeroes": (["move zeroes", "move zeros", "move all zeroes to end", "move all zeros to end of array"], []),
    "remove-duplicates": (["remove sorted duplicates"], ["remove duplicates from sorted array", "remove duplicates"]),
    "sorted-squares": (["squares in sorted order", "squares of a sorted array", "sorted squares"], []),
    "merge-sorted": (["merge sorted arrays", "merge two sorted arrays"], ["merge sorted array"]),
    "two-sum-sorted": (["two sum sorted"], ["two sum ii input array is sorted", "two sum ii"]),
    "binary-search": (["binary search", "find in a sorted array"], []),
    "search-insert": (["search insert position"], []),
    "first-occurrence": (["first occurrence", "first occurrence in sorted array"], ["find first and last position of element in sorted array", "first and last occurrences of x", "first and last occurrence"]),
    "last-occurrence": (["last occurrence", "last occurrence in sorted array"], []),
    "max-window-sum": (["maximum window sum", "max sum subarray of size k", "maximum sum subarray of size k"], []),
    "average-window": (["maximum window average", "maximum average subarray i", "maximum average subarray"], []),
    "longest-unique": (["longest unique substring", "longest substring without repeating characters"], []),
    # GfG's "Stock buy and sell" allows many transactions: a different contract under a similar name.
    "best-profit": (["best time to buy and sell", "best time to buy and sell stock"], ["stock buy and sell"]),
    "max-subarray": (["maximum subarray", "maximum subarray sum", "kadanes algorithm", "kadane algorithm", "largest sum contiguous subarray"], []),
    # The data-structure labs. "close": the row's problem differs slightly (input format, a variant of the contract).
    "reverse-list": (["reverse a linked list", "reverse linked list", "reverse a ll", "reverse a singly linked list", "reverse a linkedlist"], ["reverse a doubly linked list"]),
    "merge-two-lists": (["merge two sorted lists", "merge sorted lists", "merge two sorted linked lists"], []),
    "valid-parentheses": (["valid parentheses", "valid parenthesis", "balanced parentheses", "balanced parenthesis", "balanced paranthesis", "parenthesis checker"], []),
    "daily-temperatures": (["daily temperatures"], []),
    "recent-calls": (["number of recent calls"], []),
    "window-max": (["sliding window maximum", "maximum of all subarrays of size k"], []),
    "kth-largest": (["kth largest element in an array", "k th largest element in an array", "kth largest element"], []),
    "last-stone": (["last stone weight"], []),
    # GfG / takeUforward's "Power Set" returns a string's subsequences in sorted order, not a list's subsets.
    "subsets": (["subsets", "all subsets"], ["subsets i", "power set"]),
    "permutations": (["permutations", "all permutations"], []),
    "max-depth": (["maximum depth of binary tree", "maximum depth in bt", "max depth of binary tree"], ["height of binary tree"]),
    "level-order": (["level order traversal", "binary tree level order traversal"], []),
    "implement-trie": (["implement trie", "implement trie prefix tree", "trie implementation and operations"], ["implement trie ii prefix tree", "trie implementation and advanced operations"]),
    "prefix-counts": (["count words by prefix"], []),
    "count-components": (["connected groups"], ["number of provinces", "connected components", "number of connected components in an undirected graph"]),
    "shortest-path": (["fewest steps between nodes"], ["shortest path in undirected graph with unit weights"]),
    "count-islands": ([], ["number of islands"]),
    "shortest-grid-path": (["shortest path in a grid"], ["shortest path in binary matrix"]),
    "climb-ways": (["climbing stairs", "climb stairs"], []),
    "house-robber": (["house robber"], []),
    "single-number": (["single number", "single number i"], []),
    "count-bits": (["counting bits"], []),
}


def normal(text):
    text = str(text).lower().replace("&", " and ").replace("’", "'").replace("'", "")
    text = re.sub(r"^\s*(?:(?:lc|leetcode|q|problem)\s*)?#?\d+\s*[.):\-]\s*", "", text)  # "1. Two Sum", "LC 1 - Two Sum"
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


def key_name(key):
    """normal() of a data key: pages repeat the same few keys thousands of times. Only short keys are remembered."""
    return short_key_name(key) if len(key) <= 80 else normal(key)


@functools.lru_cache(maxsize=4096)
def short_key_name(key):
    return normal(key)


INDEX = {}
for lab, (same, close) in ALIASES.items():
    for name in same:
        INDEX[normal(name)] = (lab, "same")
    for name in close:
        INDEX.setdefault(normal(name), (lab, "close"))


def slug_of(url):
    """The problem slug of a judge link: leetcode.com/problems/two-sum → 'two sum'."""
    if not url:
        return ""
    parts = [p for p in urllib.parse.urlsplit(url).path.split("/") if p]
    host = urllib.parse.urlsplit(url).hostname or ""
    if (host == "takeuforward.org" or host.endswith(".takeuforward.org")) and len(parts) >= 3 and parts[0] == "practice":
        parts = ["practice", urllib.parse.unquote(parts[-1])]  # /practice/dsa/<slug>: the subject sits between.
    for marker in ("problems", "problem", "challenges", "practice"):
        if marker in parts and parts.index(marker) + 1 < len(parts):
            slug = parts[parts.index(marker) + 1]
            return normal(re.sub(r"(?<!\d)\d+$", "", slug.replace("-", " ").replace("_", " ")))  # GfG adds digits to slugs
    return ""


def match_row(title, *urls):
    """(lab id, 'same' | 'close') when a title or one of the row's links names a built-in lab exactly."""
    for key in (normal(title), *(slug_of(u) for u in urls)):
        if key and key in INDEX:
            return INDEX[key]
    return None, None


def same_site(url, base):
    a, b = urllib.parse.urlsplit(url).hostname or "", urllib.parse.urlsplit(base).hostname or ""
    strip = lambda h: h[4:] if h.startswith("www.") else h
    return bool(a and b) and strip(a) == strip(b)


def choose_links(candidates, base=""):
    """A row's main link and its other attached links. The sheet's own site comes first: a problem is read
    and opened where the sheet lists it; links to judges such as LeetCode or GfG are kept as alternatives."""
    found = []
    for link in candidates:
        if link and URL.match(link) and link not in found:
            found.append(link[:500])
    own = next((u for u in found if base and same_site(u, base)), "")
    judge = [u for u in found if is_judge(u) and slug_of(u)]
    if own:
        main = own
    elif not base:
        main = found[0] if found else ""  # A file has no site: its own link column comes first.
    else:
        main = judge[0] if judge else found[0] if found else ""
    return main, [u for u in judge if u != main][:5]


def title_from_url(url):
    slug = slug_of(url)
    if slug:
        return slug.title()
    parts = urllib.parse.urlsplit(url or "")
    tail = " ".join(p for p in parts.path.split("/") if p)[-60:]
    return f"{(parts.hostname or '').removeprefix('www.')} {tail}".strip() if parts.hostname else ""


# ----------------------------- reading a sheet -----------------------------
HEADERS = {
    "title": ("problem", "question", "title", "name", "task", "problems", "questions", "problem name", "question name"),
    "url": ("link", "url", "href", "problem link", "question link", "practice link", "leetcode", "gfg", "links", "problem url"),
    "difficulty": ("difficulty", "level", "diff"),
    "topic": ("topic", "category", "pattern", "tag", "tags", "section", "concept", "step", "type"),
}
# takeUforward grades its sheets basic / core / pro.
DIFFICULTY = {"easy": "Easy", "e": "Easy", "basic": "Easy", "school": "Easy", "beginner": "Easy", "core": "Medium", "pro": "Hard",
              "medium": "Medium", "m": "Medium", "med": "Medium", "hard": "Hard", "h": "Hard", "difficult": "Hard"}
NOISE = re.compile(r"^(?:true|false|yes|no|done|todo|pending|solved|unsolved|revision|revise|[✓✔✗✘x☐☑✅❌\-–—*•#.]+|\d+(?:\.\d+)?)$", re.I)


def clean_title(text):
    text = re.sub(r"\s+", " ", str(text or "")).strip()
    text = re.sub(r"^(?:(?:[-*•▪◦·>]|\[[ xX✓]?\]|[☐☑✅✓✔❌]|\d{1,4}\s*[.)\]:]|[qQ]\d+[.):])\s*)+", "", text)
    return text.strip(" -–—|:\t")[:160]


def difficulty_of(text):
    return DIFFICULTY.get(str(text or "").strip().lower(), "")


def column_roles(header):
    roles = {}
    for i, cell in enumerate(header):
        key = normal(cell)
        for role, names in HEADERS.items():
            if role not in roles and (key in names or any(key.startswith(n + " ") for n in names)):
                roles[role] = i
                break
    return roles if "title" in roles or "url" in roles else {}


def is_header(texts):
    """A header row names columns: short cells, no links, and at least two roles or one exact column name."""
    roles = column_roles(texts)
    if not roles or any(URL.search(t) or len(t) > 40 for t in texts):
        return False
    return len(roles) >= 2 or any(normal(t) in HEADERS[role] for role in roles for t in [texts[roles[role]]])


def rows_from_table(table, default_topic="", base=""):
    """Rows of cells (with optional hyperlinks) → sheet rows. A header row assigns roles; without
    one, the first text cell is the title and any link cell the link. A text-only row in a sheet of
    links is a section heading and becomes the topic of the rows below it."""
    table = [[(str(c[0]).strip() if c[0] is not None else "", c[1]) for c in row] for row in table]
    table = [row for row in table if any(text or link for text, link in row)]
    if not table:
        return []
    header_at = next((i for i, row in enumerate(table[:8]) if is_header([text for text, _ in row])), None)
    roles = column_roles([text for text, _ in table[header_at]]) if header_at is not None else {}
    body = table[header_at + 1:] if header_at is not None else table
    linked = sum(1 for row in body if any(link or URL.search(text) for text, link in row))
    rows, topic = [], default_topic
    for row in body:
        texts = [text for text, _ in row]
        link = next((link for _, link in row if link), "") or next((URL.search(t).group(0) for t in texts if URL.search(t)), "")
        if roles:
            pick = lambda role: texts[roles[role]] if role in roles and roles[role] < len(texts) else ""
            title = clean_title(pick("title")) if "title" in roles else ""
            if "url" in roles and URL.search(pick("url") or ""):
                link = URL.search(pick("url")).group(0)
            elif "url" in roles and roles["url"] < len(row) and row[roles["url"]][1]:
                link = row[roles["url"]][1]
            if "title" in roles and roles["title"] < len(row) and row[roles["title"]][1]:
                link = link or row[roles["title"]][1]
            level, row_topic = difficulty_of(pick("difficulty")), clean_title(pick("topic"))
        else:
            title = next((clean_title(t) for t in texts if t and not URL.fullmatch(t) and not NOISE.match(t.strip()) and not difficulty_of(t) and len(clean_title(t)) > 1), "")
            level, row_topic = next((difficulty_of(t) for t in texts if difficulty_of(t)), ""), ""
        title = title if title and not NOISE.match(title) else ""
        if not link and title and linked >= max(2, len(body) // 3) and sum(1 for t in texts if t) == 1:
            topic = title  # A heading row inside a sheet of links.
            continue
        main, others = choose_links([link] + [l for _, l in row if l] + [m.group(0) for t in texts for m in URL.finditer(t)], base)
        title = title or title_from_url(main)
        if title:
            rows.append(dict(title=title, url=main, links=others, difficulty=level, topic=(row_topic or topic)[:80]))
    return rows


# Starts only where a run of spaces starts, and never gives tokens back: a long gap or a run of status words stays linear.
STATUS_TAIL = re.compile(r"(?<!\s)(?:\s+(?:done|todo|to do|solved|unsolved|pending|revise|revision|yes|no|true|false|[✓✔✗✘☐☑✅❌xX]|\d{1,2}/\d{1,2}(?:/\d{2,4})?))++\s*$", re.I)
HEADER_WORDS = {"s", "no", "sno", "sr", "problem", "problems", "question", "questions", "title", "name", "link", "links", "url", "difficulty", "level",
                "topic", "topics", "status", "done", "practice", "solution", "notes", "category", "pattern", "tags", "day", "date", "revision"}


def ocr_line(raw):
    """Text recognised from a photo of a table: drop serial numbers, status columns and a header row."""
    if set(normal(raw).split()) <= HEADER_WORDS:
        return ""
    raw = re.sub(r"^\s*\d{1,4}(?:\s*[.)|:])?\s++(?=\S*[A-Za-z])", "", raw)
    level = re.search(r"\b(easy|medium|hard)\b", raw, re.I)
    if level:  # The difficulty column ends the useful part of a row; status and notes follow it.
        raw = raw[:level.start()].rstrip(" |") + " - " + level.group(1)
    return STATUS_TAIL.sub("", raw.replace("|", " ")).strip()


def read_photo_text(text):
    """(rows, title) from text recognised in a picture: a line above the table's header row is the
    sheet's title rather than a problem."""
    lines = [line for line in text.replace("\r", "\n").split("\n") if line.strip()]
    header = next((i for i, line in enumerate(lines[:12]) if not ocr_line(line.strip()) and len(normal(line).split()) >= 2), None)
    title = clean_title(lines[0]) if header else ""
    return rows_from_text("\n".join(lines[header + 1:] if header is not None else lines), ocr=True), title[:80]


def md_link(text):
    """The first Markdown link [title](http…) in text, as re.search(r"\\[([^\\]]+)\\]\\((https?://[^)]+)\\)") finds it:
    (start, end, title, url) or None. Read bracket by bracket, so text full of [ or ]( costs one pass, not one per mark."""
    after, paren = 0, -1
    while True:
        close = text.find("]", after)
        if close < 0:
            return None
        start = text.find("[", after, close)  # Any later [ before this ] would end at the same ].
        if 0 <= start < close - 1 and text.startswith(("(http://", "(https://"), close + 1):
            if paren <= close:
                paren = text.find(")", close + 1)
                if paren < 0:
                    return None  # No ) left, so no later link can end either.
            url = text[close + 2:paren]
            if len(url) > len("https://" if url.startswith("https://") else "http://"):
                return start, paren + 1, text[start + 1:close], url
        after = close + 1


def rows_from_text(text, ocr=False):
    """Lines of a list, a markdown document, or text recognised from an image."""
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if sum(1 for line in lines if line.strip().startswith("|")) >= 3:  # A markdown table.
        table = [[(cell.strip(), "") for cell in line.strip().strip("|").split("|")] for line in lines if line.strip().startswith("|") and not re.fullmatch(r"[\s|:\-]+", line)]
        for row in table:
            for i, (cell, _) in enumerate(row):
                found = md_link(cell)
                if found:
                    row[i] = found[2:]
        return rows_from_table(table)
    rows, topic = [], ""
    for line in lines:
        raw = ocr_line(line.strip()) if ocr else line.strip()
        if not raw:
            continue
        if re.match(r"^\[[^\]]+\]:\s*\S+", raw) or raw.startswith(("```", "<!--", "![")):
            continue  # Markdown link definitions, code fences, comments and images.
        heading = re.match(r"^(?:#{1,6}\s+(.+)|(.{2,60}):\s*$)", raw)
        if heading and not URL.search(raw):
            topic = clean_title(heading.group(1) or heading.group(2))
            continue
        link = ""
        found = md_link(raw)
        if found:
            raw, link = raw.replace(raw[found[0]:found[1]], found[2]), found[3]
        elif URL.search(raw):
            link = URL.search(raw).group(0)
            raw = raw.replace(link, " ")
        level = ""
        tag = re.search(r"(?<![\s(\[|,–-])[\s(\[|,–-]+(easy|medium|hard)[\s)\]|,]*$", raw, re.I)
        if tag:
            level, raw = DIFFICULTY[tag.group(1).lower()], raw[:tag.start()]
        title = clean_title(raw)
        if (not title or NOISE.match(title) or len(title) < 3) and link:
            title = title_from_url(link)
        words = len(title.split())
        prose = not link and (words > 14 or (words > 7 and title.endswith((".", ",", ";"))))  # A sentence, not a problem name.
        if title and not prose and not NOISE.match(title) and len(title) >= 3 and re.search(r"[A-Za-z]{2}", title):
            rows.append(dict(title=title, url=link[:500], difficulty=level, topic=topic[:80]))
    return rows


def rows_from_json(data):
    items = data.get("problems") or data.get("rows") or data.get("questions") if isinstance(data, dict) else data
    if not isinstance(items, list):
        raise ValueError("A JSON sheet is a list of problems, or an object with a \"problems\" list.")
    rows = []
    for item in items:
        if isinstance(item, str):
            rows.extend(rows_from_text(item))
        elif isinstance(item, dict):
            lower = {normal(k): v for k, v in item.items()}
            title = next((lower[k] for k in HEADERS["title"] if isinstance(lower.get(k), str)), "")
            link = next((lower[k] for k in HEADERS["url"] if isinstance(lower.get(k), str) and URL.match(lower[k])), "")
            level = next((difficulty_of(lower[k]) for k in HEADERS["difficulty"] if isinstance(lower.get(k), str)), "")
            topic = next((lower[k] for k in HEADERS["topic"] if isinstance(lower.get(k), str)), "")
            title = clean_title(title) or title_from_url(link)
            if title:
                rows.append(dict(title=title, url=link[:500], difficulty=level, topic=clean_title(topic)[:80]))
    return rows


def rows_from_xlsx(blob):
    try:
        import openpyxl
    except ImportError:
        raise ValueError("Reading Excel files needs openpyxl (pip install openpyxl). You can also save the sheet as CSV and upload that.")
    unreadable = "This Excel file could not be read. Save it as .xlsx or CSV and try again."
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            members = archive.infolist()
    except Exception:
        raise ValueError(unreadable)
    # A zip never unpacks past its stated sizes, so they bound the work before anything is parsed.
    if len(members) > 2000 or sum(m.file_size for m in members) > MAX_XLSX:
        raise ValueError(f"This Excel file unpacks to more than {MAX_XLSX // 1_000_000} MB, far more than a problem sheet needs. Save the sheet as CSV and upload that.")
    try:
        book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)  # Streams rows instead of building every cell.
    except Exception:
        raise ValueError(unreadable)
    rows = []
    try:
        for sheet in book.worksheets:  # Many sheets keep one tab per topic: the tab name is the topic.
            links = sheet_links(book, sheet)
            table = [[(cell.value, links.get((r, c), "")) for c, cell in enumerate(row, 1)] for r, row in enumerate(sheet.iter_rows(max_row=2000, max_col=30), 1)]
            topic = "" if len(book.worksheets) == 1 or re.fullmatch(r"sheet\s*\d*", sheet.title, re.I) else sheet.title
            rows.extend(rows_from_table(table, topic))
    except Exception:  # A read-only book parses each sheet as it is read, so a broken sheet fails here.
        raise ValueError(unreadable)
    finally:
        book.close()
    return rows


def sheet_links(book, sheet):
    """A read-only sheet's hyperlinks by (row, column). openpyxl binds them only when it loads every cell, so they
    are read from the sheet's own XML and relationships, as openpyxl does then."""
    from openpyxl.packaging.relationship import get_dependents, get_rels_path
    from openpyxl.utils.cell import range_boundaries
    from openpyxl.xml.constants import REL_NS, SHEET_MAIN_NS
    from openpyxl.xml.functions import iterparse
    archive, path = book._archive, sheet._worksheet_path
    rels_path = get_rels_path(path)
    targets = {r.Id: r.Target for r in get_dependents(archive, rels_path) if r.Target} if rels_path in archive.namelist() else {}
    links = {}
    if not targets:
        return links  # Links to other places in the workbook have no target to keep.
    with archive.open(path) as source:
        for _, node in iterparse(source):
            target = targets.get(node.get(f"{{{REL_NS}}}id")) if node.tag == f"{{{SHEET_MAIN_NS}}}hyperlink" else None
            try:
                bounds = range_boundaries(node.get("ref") or "") if target else None
            except ValueError:
                bounds = None  # A link without a readable cell reference has no cell to attach to.
            if bounds:  # "A2", "A2:B3", or a whole column or row ("A:A", "2:2").
                low_col, low_row, high_col, high_row = (b if b is not None else d for b, d in zip(bounds, (1, 1, 30, 2000)))
                for r in range(low_row, min(high_row, 2000) + 1):
                    for c in range(low_col, min(high_col, 30) + 1):
                        links[(r, c)] = target
            node.clear()
    return links


def rows_from_csv(text):
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel_tab if sample.count("\t") > sample.count(",") else csv.excel
    try:
        table = [[(cell, "") for cell in row] for row in csv.reader(io.StringIO(text), dialect)]
    except csv.Error as error:
        if "field limit" in str(error):
            raise ValueError(f"A cell in this sheet is longer than {csv.field_size_limit():,} characters, so it can't be read as a table. Check the file, or paste the problem titles instead.")
        raise ValueError(f"This sheet couldn't be read as CSV ({error}). Check the file, or paste the problem titles instead.")
    if max((len(row) for row in table), default=0) <= 1:
        return rows_from_text(text)
    for row in table:
        for i, (cell, _) in enumerate(row):
            formula = re.match(r'=HYPERLINK\("([^"]+)"\s*[,;]\s*"([^"]*)"\)', cell, re.I)
            if formula:
                row[i] = (formula.group(2), formula.group(1))
    return rows_from_table(table)


class PageLinks(html.parser.HTMLParser):
    """Tables, headings and links of a web page, as plain data. Relative links resolve against the page."""
    def __init__(self, base=""):
        super().__init__(convert_charrefs=True)
        self.base = base
        self.tables, self.links, self.title = [], [], ""
        self._table = self._row = self._cell = None
        self._href, self._text, self._heading, self._heading_text, self._in_title, self._skip = None, [], "", "", False, 0

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in ("script", "style", "noscript", "svg"):
            self._skip += 1
        elif tag == "table":
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = [[], ""]
        elif tag == "a":
            self._href, self._text = attrs.get("href") or "", []
        elif tag in ("h1", "h2", "h3", "h4"):
            self._heading = tag
            self._text = []
        elif tag == "title":
            self._in_title = True

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript", "svg"):
            self._skip = max(0, self._skip - 1)
        elif tag == "a" and self._href is not None:
            text = clean_title(" ".join(self._text))
            if self.base and self._href.startswith("/") and not self._href.startswith("//"):
                self._href = urllib.parse.urljoin(self.base, self._href)
            if self._href.startswith("http"):
                self.links.append((text, self._href, self._heading_text))
                if self._cell is not None and not self._cell[1]:
                    self._cell[1] = self._href
            self._href = None
        elif tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append((" ".join(" ".join(self._cell[0]).split()), self._cell[1]))
            self._cell = None
        elif tag == "tr" and self._row is not None and self._table is not None:
            self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            self.tables.append(self._table)
            self._table = None
        elif tag == self._heading:
            self._heading_text, self._heading = clean_title(" ".join(self._text)), ""
        elif tag == "title":
            self._in_title = False

    def handle_data(self, data):
        if self._skip:
            return
        if self._in_title:
            self.title += data
        if self._cell is not None:
            self._cell[0].append(data)
        if self._href is not None or self._heading:
            self._text.append(data)


# ----------------------------- data a page embeds -----------------------------
# Many sheet sites render their list in the browser from data embedded in the page: Next.js page
# data, React server-component payloads, or JSON script tags. That data is read as plain JSON.
TITLE_KEYS = {"title", "label", "name", "problem", "question", "problem name", "problemname", "problem title", "problemtitle",
              "question name", "questionname", "question title", "questiontitle", "display name", "displayname"}
CHILD_KEYS = {"children", "items", "problems", "questions", "subsections", "sections", "topics", "steps", "lectures", "subtopics",
              "modules", "chapters", "categories", "entries", "nodes", "list"}
KIND_KEYS = {"type", "layouttype", "kind", "contenttype", "itemtype", "item type", "content type", "category type"}
NOT_PROBLEMS = {"learning", "article", "lesson", "video", "contest", "blog", "theory", "note", "notes", "quiz", "editorial", "reading", "course", "mcq"}
PROBLEM_KINDS = {"practice", "problem", "question", "exercise", "challenge", "coding"}


SCRIPT_OPEN, SCRIPT_CLOSE = re.compile(r"<script", re.I), re.compile(r"</script>", re.I)
JSON_TYPE = re.compile(r"type=[\"']application/(?:ld\+)?json[\"']", re.I)


def json_scripts(page):
    """The bodies of a page's JSON script tags. Found tag by tag rather than by one pattern, so a page full of
    unclosed tags costs one pass instead of a scan to the end of the page for each tag."""
    bodies, closing, after = [], -1, 0
    for opening in SCRIPT_OPEN.finditer(page):
        if opening.start() < after:
            continue  # Inside the body of a script already read.
        if closing < opening.end():
            closing = page.find(">", opening.end())
        if closing < 0:
            break
        if not JSON_TYPE.search(page, opening.end(), closing):
            continue
        end = SCRIPT_CLOSE.search(page, closing + 1)
        if not end:
            break  # No later tag can close either.
        bodies.append(page[closing + 1:end.start()])
        after = end.end()
    return bodies


def embedded_json(page):
    """JSON values a page embeds for its own scripts. Data nested too deeply to read is skipped like any unreadable data."""
    values = []
    for body in json_scripts(page):
        try:
            values.append(json.loads(body))
        except (ValueError, RecursionError):
            pass
    pushes = re.findall(r'self\.__next_f\.push\(\[\d+,("(?:[^"\\]|\\.)*+")\]\)', page)
    try:
        flight = "".join(json.loads(chunk) for chunk in pushes)
    except (ValueError, RecursionError):
        flight = ""
    for line in flight.split("\n"):  # Server-component payload: one "id:JSON" record per line.
        found = re.match(r"[0-9a-f]+:([\[{].*)", line, re.S)
        if found:
            try:
                values.append(json.loads(found.group(1)))
            except (ValueError, RecursionError):
                pass
    return values


MAX_BUILT = 20_000  # Objects one schema-indexed table may expand into; a real sheet needs a few thousand.
MAX_VISITS = 300_000  # Objects and lists one page's data may visit while looking for problems.


class TooLarge(Exception):
    pass


def expand_columns(value, depth=0):
    """Expand schema-indexed tables ({"fields": [[names…]…], "rows": [[schema, values…]…], "roots": […]})
    into ordinary objects, resolving child row indices into nested objects. A row is never nested inside itself,
    and a table that would expand past MAX_BUILT objects (rows shared by many parents) is skipped, not cut short."""
    if depth > 80:
        return None
    if isinstance(value, list):
        return [expand_columns(v, depth + 1) for v in value]
    if not isinstance(value, dict):
        return value
    schemas, rows = value.get("fields", value.get("columns")), value.get("rows")
    if (isinstance(schemas, list) and schemas and all(isinstance(f, list) and all(isinstance(k, str) for k in f) for f in schemas)
            and isinstance(rows, list) and rows and all(isinstance(r, list) and r and type(r[0]) is int and 0 <= r[0] < len(schemas) for r in rows)):
        parents, built = [], [0]

        def build(index, level=0):
            built[0] += 1
            if built[0] > MAX_BUILT:
                raise TooLarge
            row = rows[index]
            item = dict(zip(schemas[row[0]], row[1:]))
            parents.append(index)
            for key, child in item.items():
                if key_name(key) in CHILD_KEYS and isinstance(child, list) and all(type(c) is int and 0 <= c < len(rows) for c in child):
                    item[key] = [build(c, level + 1) for c in child if c not in parents] if level < 12 else []
            parents.pop()
            return item
        roots = value.get("roots")
        try:
            if isinstance(roots, list) and roots and all(type(r) is int and 0 <= r < len(rows) for r in roots):
                return [build(r) for r in roots]
            return [build(i) for i in range(len(rows))]
        except TooLarge:
            return None
    return {k: expand_columns(v, depth + 1) for k, v in value.items()}


def is_judge(url):
    host = urllib.parse.urlsplit(url).hostname or ""
    return any(host == j or host.endswith("." + j) for j in JUDGES)


def site_page(node, base):
    """The problem's own page on the sheet's site, when the site's data names it. takeUforward's rows carry
    redirectTo {layoutType, contentType, itemSlug}; its scripts route that to /{layoutType}/{contentType}/{itemSlug}."""
    host = urllib.parse.urlsplit(base).hostname or ""
    target = node.get("redirectTo")
    if (host != "takeuforward.org" and not host.endswith(".takeuforward.org")) or not isinstance(target, dict):
        return ""
    layout, content, slug = target.get("layoutType"), target.get("contentType") or "dsa", target.get("itemSlug")
    if layout != "practice" or not isinstance(content, str) or not re.fullmatch(r"[a-z-]{1,30}", content) or not isinstance(slug, str) or not slug.strip():
        return ""
    quoted = urllib.parse.quote(slug.strip(), safe="'()*!~")  # As the site's encodeURIComponent writes it.
    return f"https://takeuforward.org/{layout}/{content}/{quoted}"[:500]


def json_rows(values, base=""):
    """Problems inside embedded data: an object with a title and evidence of being a problem (a link to a
    coding judge, a difficulty, or a practice kind). Objects with a title and a list of children are
    sections; their titles become the topic. Lessons, articles and contests are skipped."""
    rows, visits = [], [0]

    def title_of(node):
        for key, value in node.items():
            if key_name(key) in TITLE_KEYS and isinstance(value, str) and 2 <= len(value.strip()) <= 160 and not URL.match(value.strip()):
                return clean_title(value)
        return ""

    def links_of(node):
        found = []
        for key, value in node.items():
            name = key_name(key)
            if any(word in name for word in ("blog", "video", "yt", "editorial", "solution", "article", "image", "icon", "avatar", "thumbnail")):
                continue  # Lessons, videos and pictures are not the problem's page.
            if isinstance(value, str) and URL.fullmatch(value.strip()):
                found.append(value.strip())
            elif isinstance(value, str) and base and value.startswith("/") and not value.startswith("//") and any(w in name for w in ("url", "link", "href", "path")):
                found.append(urllib.parse.urljoin(base, value.strip()))
            elif isinstance(value, dict):
                found.extend(v.strip() for v in value.values() if isinstance(v, str) and URL.fullmatch(v.strip()))
        return found

    def visit(node, trail, depth):
        if depth > 80 or not isinstance(node, (list, dict)):
            return
        visits[0] += 1
        if visits[0] > MAX_VISITS:
            return  # However the data repeats itself, reading it stays bounded.
        if isinstance(node, list):
            for item in node:
                visit(item, trail, depth + 1)
            return
        title = title_of(node)
        children = [v for k, v in node.items() if key_name(k) in CHILD_KEYS and isinstance(v, list) and any(isinstance(c, dict) for c in v)]
        if title and children:
            for child in children:
                visit(child, trail + [title], depth + 1)
            return
        kinds = {normal(v) for k, v in node.items() if key_name(k) in KIND_KEYS and isinstance(v, str)}
        links = links_of(node)
        judge = next((u for u in links if is_judge(u) and slug_of(u)), "")
        level = next((difficulty_of(v) for k, v in node.items() if key_name(k) in HEADERS["difficulty"] and isinstance(v, str) and difficulty_of(v)), "")
        if title and not kinds & NOT_PROBLEMS and (judge or level or kinds & PROBLEM_KINDS):
            source = site_page(node, base)
            main, others = choose_links([source] + links, base)
            rows.append(dict(title=title, url=main, links=others, source=source or (main if base and same_site(main, base) else ""), difficulty=level, topic=" · ".join(trail[-2:])[:80]))
            return
        for value in node.values():
            if isinstance(value, (dict, list)):
                visit(value, trail, depth + 1)

    for value in values:
        visit(expand_columns(value), [], 0)
    return rows


def rows_from_html(text, base=""):
    page = PageLinks(base)
    page.feed(text)
    rows = []
    for table in page.tables:
        if sum(1 for row in table if any(link for _, link in row)) >= 3:
            rows.extend(rows_from_table(table, base=base))
    embedded = json_rows(embedded_json(text), base)
    if len(embedded) >= 3 and len(embedded) > len(rows):  # The page's own data is the fuller list.
        rows = embedded
    if not rows:  # No table of links: take links to coding judges, in page order.
        seen = set()
        for text, href, heading in page.links:
            if is_judge(href) and slug_of(href) and href not in seen:
                seen.add(href)
                rows.append(dict(title=text if len(text) >= 3 and not NOISE.match(text) else title_from_url(href), url=href[:500], links=[], difficulty="", topic=heading[:80]))
    return rows, clean_title(re.split(r"(?<!\s)\s+[|–—-]\s+", page.title.strip())[0])[:80]


def read_upload(filename, blob):
    """(rows, name) from an uploaded file's bytes."""
    name = re.sub(r"\.[A-Za-z0-9]+$", "", filename or "My sheet")[:80] or "My sheet"
    ext = (filename or "").lower().rsplit(".", 1)[-1] if "." in (filename or "") else ""
    if ext in ("xlsx", "xlsm"):
        return rows_from_xlsx(blob), name
    if ext == "xls":
        raise ValueError("Old .xls files can't be read. Save the sheet as .xlsx or CSV and upload that.")
    if ext in ("png", "jpg", "jpeg", "webp", "gif", "bmp", "pdf"):
        raise ValueError("Use the Image option for a photo or screenshot of your sheet.")
    text = blob.decode("utf-8-sig", errors="replace")
    if ext == "json":
        try:
            return rows_from_json(json.loads(text)), name
        except (json.JSONDecodeError, RecursionError):
            raise ValueError("This JSON file could not be read.")
    if ext in ("html", "htm"):
        return rows_from_html(text)[0], name
    if ext in ("csv", "tsv"):
        return rows_from_csv(text), name
    return rows_from_text(text), name


# ----------------------------- reading a link -----------------------------
def sheet_export_url(url):
    """A Google Sheets link → its CSV export for the same tab."""
    parts = urllib.parse.urlsplit(url)
    found = re.match(r"/spreadsheets/d/(?:e/)?([A-Za-z0-9_-]+)", parts.path)
    if parts.hostname != "docs.google.com" or not found:
        return None
    gid = re.search(r"gid=(\d+)", parts.fragment + "&" + parts.query)
    if "/d/e/" in parts.path:  # A "Publish to the web" link.
        return f"https://docs.google.com/spreadsheets/d/e/{found.group(1)}/pub?output=csv" + (f"&gid={gid.group(1)}" if gid else "")
    return f"https://docs.google.com/spreadsheets/d/{found.group(1)}/export?format=csv" + (f"&gid={gid.group(1)}" if gid else "")


NAT64, LOCAL_NAT64, IPV4_COMPATIBLE = (ipaddress.ip_network(n) for n in ("64:ff9b::/96", "64:ff9b:1::/48", "::/96"))


def is_public(address):
    """A public internet address, including the IPv4 address an IPv6 one carries (::ffff:a.b.c.d, ::a.b.c.d,
    NAT64 64:ff9b::/96, 6to4, Teredo): ::127.0.0.1 reaches this machine just as 127.0.0.1 does."""
    if not address.is_global:
        return False
    if address.version == 6:
        if address in LOCAL_NAT64:
            return False
        inner = [address.ipv4_mapped, address.sixtofour, *(address.teredo or ())]
        if address in NAT64 or address in IPV4_COMPATIBLE:
            inner.append(ipaddress.IPv4Address(int(address) & 0xFFFFFFFF))
        return all(a.is_global for a in inner if a is not None)
    return True


def public_addresses(host, port=None):
    """The addresses a host resolves to, all of them public: a sheet link must never reach this machine or its network."""
    try:
        infos = socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
    except (socket.gaierror, UnicodeError):  # UnicodeError: a name no DNS label can hold.
        raise ValueError(f"Couldn't find {host[:80]}. Check the link and your connection.")
    for info in infos:
        if not is_public(ipaddress.ip_address(info[4][0].split("%")[0])):
            raise ValueError("Links to private or local addresses can't be imported.")
    return infos


def public_host(host):
    public_addresses(host)


class FetchError(ValueError):
    """A fetch that failed, with the site's HTTP status when it answered with one."""
    def __init__(self, message, status=None):
        super().__init__(message)
        self.status = status


class DeadlineReads:
    """A socket whose reads stop at a deadline for the whole fetch, not only after each quiet spell: a site that
    drips one byte every few seconds can't hold a request open for hours."""
    deadline = None

    def recv_into(self, buffer, *args):
        if self.deadline is not None:
            left = self.deadline - time.monotonic()
            if left <= 0:
                raise TimeoutError("fetch deadline passed")
            self.settimeout(min(left, 10))
        return super().recv_into(buffer, *args)


class DeadlineSocket(DeadlineReads, socket.socket):
    pass


class DeadlineSSLSocket(DeadlineReads, ssl.SSLSocket):
    pass


def connect_public(address, timeout, source_address=None, *_, deadline=None, check=True):
    """socket.create_connection, connecting only to the addresses just checked: the name is resolved once, so a
    second lookup can't point the connection somewhere private (DNS rebinding)."""
    host, port = address
    infos = public_addresses(host, port) if check else socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
    timeout, failure = timeout if isinstance(timeout, (int, float)) else 10, None
    for family, kind, proto, _, target in infos:
        sock = DeadlineSocket(family, kind, proto)
        try:
            sock.settimeout(min(timeout, max(deadline - time.monotonic(), 0.001)) if deadline else timeout)
            if source_address:
                sock.bind(source_address)
            sock.connect(target)
            sock.deadline = deadline
            return sock
        except OSError as error:
            failure = error
            sock.close()
    raise failure or OSError(f"Couldn't connect to {host}.")


class PublicHTTP(http.client.HTTPConnection):
    def __init__(self, *args, deadline=None, check=True, **kwargs):
        super().__init__(*args, **kwargs)
        self._create_connection = functools.partial(connect_public, deadline=deadline, check=check)


class PublicHTTPS(http.client.HTTPSConnection):
    """HTTPS to the checked address. TLS still names the link's host, so SNI and the certificate check are unchanged."""
    def __init__(self, *args, deadline=None, check=True, **kwargs):
        context = ssl.create_default_context()  # Certificates and host names verified, as urllib's default context does.
        context.set_alpn_protocols(["http/1.1"])
        context.sslsocket_class = DeadlineSSLSocket
        super().__init__(*args, **{**kwargs, "context": context})
        self._create_connection = functools.partial(connect_public, deadline=deadline, check=check)
        self.deadline = deadline

    def connect(self):
        super().connect()
        self.sock.deadline = self.deadline


def proxied(req):
    return bool(req.has_proxy() or getattr(req, "_tunnel_host", None))


class PublicHandler(urllib.request.HTTPHandler, urllib.request.HTTPSHandler):
    """Plain and TLS connections that keep the fetch's deadline and connect only to checked addresses. Through a
    proxy the proxy resolves the site, so only the link's own check (before the fetch) applies."""
    def __init__(self, deadline):
        urllib.request.AbstractHTTPHandler.__init__(self)  # HTTPSHandler's own context is never used: each connection makes one.
        self.deadline = deadline

    def http_open(self, req):
        return self.do_open(functools.partial(PublicHTTP, deadline=self.deadline, check=not proxied(req)), req)

    def https_open(self, req):
        return self.do_open(functools.partial(PublicHTTPS, deadline=self.deadline, check=not proxied(req)), req)


# While a problem page is read, each redirect must also be one robots.txt allows (see read_problem).
REDIRECT_CHECK = contextvars.ContextVar("redirect_check", default=None)


class CheckedRedirects(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        parts = urllib.parse.urlsplit(newurl)
        if parts.scheme not in ("http", "https") or not parts.hostname:
            raise ValueError("The link redirected somewhere that can't be imported.")
        if is_leetcode(newurl):
            raise ValueError("The link redirected to LeetCode, which builds its pages in the browser, so it isn't read here.")
        public_host(parts.hostname)
        check = REDIRECT_CHECK.get()
        if check and not check(newurl):
            raise ValueError(f"The link redirected to {parts.hostname}, which asks automated tools not to read that page (robots.txt).")
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def fetch(url):
    """(body bytes, final URL, content type, charset), refusing private addresses and large or slow responses."""
    parts = urllib.parse.urlsplit(url)
    if parts.scheme not in ("http", "https") or not parts.hostname:
        raise ValueError("Paste a full link that starts with https://")
    public_host(parts.hostname)
    opener = urllib.request.build_opener(CheckedRedirects, PublicHandler(time.monotonic() + FETCH_SECONDS))
    request = urllib.request.Request(url, headers={"User-Agent": AGENT, "Accept": "text/csv,text/html,text/plain,application/json;q=0.9,*/*;q=0.5"})
    slow = FetchError("The site took too long to send the page. Try again later, or download the sheet and upload the file.")
    try:
        with opener.open(request, timeout=10) as response:
            body = response.read(MAX_FETCH + 1)
            final, kind = response.geturl(), response.headers.get_content_type()
            charset = response.headers.get_content_charset() or "utf-8"
    except ValueError:
        raise
    except urllib.error.HTTPError as error:
        if error.code in (404, 410):
            raise FetchError(f"That page wasn't found ({error.code}). Check the link.", error.code)
        if error.code in (401, 403):
            raise FetchError(f"The site refused to share that page ({error.code}). If the sheet is private, share it publicly, or download it and upload the file.", error.code)
        raise FetchError(f"The site answered with an error ({error.code}). Try again later, or download the sheet and upload the file.", error.code)
    except TimeoutError:
        raise slow
    except Exception as error:
        if isinstance(getattr(error, "reason", None), TimeoutError):
            raise slow
        raise FetchError("The link couldn't be reached. Check it, or download the sheet and upload the file instead.")
    if len(body) > MAX_FETCH:
        raise FetchError("That page is larger than 3 MB. Download the sheet and upload the file instead.")
    return body, final, kind, charset


def decode(body, charset):
    """A response's text. A charset Python doesn't know (MySQL's "utf8mb4" is UTF-8 by another name) or can't
    decode leniently is read as UTF-8 instead of failing the whole read."""
    try:
        return body.decode(charset, errors="replace")
    except (LookupError, UnicodeError):
        return body.decode("utf-8", errors="replace")


def unknown_charset(charset):
    """A charset decode() had to replace with UTF-8 (other than MySQL's names for UTF-8 itself)."""
    try:
        b"a".decode(charset, errors="replace")  # Not b"": decoding nothing never looks the codec up.
        return False
    except (LookupError, UnicodeError):
        return re.sub(r"[-_]", "", charset.lower()) not in ("utf8mb4", "utf8mb3")


def looks_like_problem(url):
    """A URL that names one problem (…/problems/<slug>, …/practice/<subject>/<slug>), not a sheet or a list."""
    path = urllib.parse.urlsplit(url).path.lower()
    if any(word in path for word in ("sheet", "list", "prep-hub", "study-plan", "studyplan", "collection", "problemset")):
        return False
    return bool(re.search(r"/(problems?|practice|challenges?|questions?|tasks?)/[^/]+", path))


def read_link(url, fetcher=fetch):
    """(rows, name) for a sheet link: Google Sheets, a CSV/JSON/text file, or a web page that lists problems."""
    url = url.strip()
    export = sheet_export_url(url)
    given = urllib.parse.urlsplit(url).hostname or ""
    if not export and slug_of(url) and any(given == j or given.endswith("." + j) for j in JUDGES):
        # A single problem's page: the link itself is the row. Its title is the page's own when the page can be read
        # (never LeetCode's, which is built in the browser); otherwise the link's words name it.
        row = dict(title=title_from_url(url), url=url[:500], difficulty="", topic="")
        if not is_leetcode(url) and looks_like_problem(url):
            try:
                problem = read_problem({"source": url}, fetcher)
            except ValueError:
                problem = None
            if problem:
                return [dict(row, title=problem["title"], source=problem["source"][:500])], problem["title"]
        return [row], "My sheet"
    if is_leetcode(url):  # LeetCode is never fetched: its lists, like its problems, are built in the browser.
        raise ValueError("LeetCode builds its lists in the browser, so they aren't read here. Paste the problem links, or take a screenshot of the list and import that.")
    body, final, kind, charset = fetcher(export or url)
    host = urllib.parse.urlsplit(final).hostname or ""
    if export and (host == "accounts.google.com" or kind == "text/html"):
        raise ValueError("This Google Sheet isn't public. In Google Sheets choose Share → General access → Anyone with the link, then paste the link again. Or use File → Download → CSV and upload the file.")
    text = decode(body, charset)
    path = urllib.parse.urlsplit(final).path.lower()
    name = urllib.parse.unquote(path.rstrip("/").rsplit("/", 1)[-1] or host)[:80]
    if export or kind in ("text/csv", "text/tab-separated-values") or path.endswith((".csv", ".tsv")):
        return rows_from_csv(text), "My Google Sheet" if export else re.sub(r"\.[a-z]+$", "", name)
    if kind == "application/json" or path.endswith(".json"):
        try:
            return rows_from_json(json.loads(text)), re.sub(r"\.[a-z]+$", "", name)
        except (json.JSONDecodeError, RecursionError):
            raise ValueError("The link returned JSON that couldn't be read.")
    if kind == "text/html":
        if looks_like_problem(final):
            # A page about one problem (its statement and examples), not a list: it becomes a one-problem sheet.
            try:
                problem = problem_from_page(text, final)
            except ValueError:
                problem = None
            if problem and (problem["examples"] or len(problem["description"]) > 80):
                return [dict(title=problem["title"], url=final[:500], source=final[:500], difficulty="", topic="")], problem["title"]
        rows, title = rows_from_html(text, final)
        if not rows:
            raise ValueError("No problem list was found on that page. Some sites load their list only after the page opens in a browser: take a screenshot of the list (or download it as a file) and import that instead.")
        return rows, title or host
    return rows_from_text(text), re.sub(r"\.[a-z]+$", "", name) or host


def finish(rows):
    """De-duplicate, cap and match rows to built-in labs."""
    seen, result = set(), []
    for row in rows:
        key = (normal(row["title"]), row.get("url", ""))
        if key in seen:
            continue
        seen.add(key)
        links = [l[:500] for l in row.get("links") or [] if l and l != row.get("url")][:5]
        lab, fit = match_row(row["title"], row.get("url", ""), *links)
        result.append(dict(title=row["title"][:160], url=row.get("url", "")[:500], links=links, source=row.get("source", "")[:500], difficulty=row.get("difficulty", ""), topic=row.get("topic", "")[:80], match=lab, fit=fit, lab=None))
        if len(result) >= MAX_ROWS:
            break
    return result


def clean_rows(rows, known_labs):
    """Rows sent back by the learner after review: plain fields only; matches are re-derived or
    must name a real built-in lab when chosen by hand."""
    if not isinstance(rows, list) or not 1 <= len(rows) <= MAX_ROWS:
        raise ValueError(f"A sheet has between 1 and {MAX_ROWS} problems.")
    result = []
    for row in rows:
        if not isinstance(row, dict) or not isinstance(row.get("title"), str) or not clean_title(row["title"]):
            raise ValueError("Every problem needs a title.")
        url = row.get("url") or ""
        if not isinstance(url, str) or (url and not re.match(r"https?://", url)):
            raise ValueError(f"“{row['title'][:40]}” has a link that doesn't start with http(s)://.")
        chosen = row.get("match")
        if chosen is not None and chosen not in known_labs:
            raise ValueError("Choose a lab from the library.")
        extra = row.get("links") if isinstance(row.get("links"), list) else []
        links = [l[:500] for l in extra if isinstance(l, str) and re.match(r"https?://", l) and l != url][:5]
        lab, fit = match_row(row["title"], url, *links)
        fit = fit if chosen == lab and chosen else "manual" if chosen else None
        lab_id = row.get("lab") if isinstance(row.get("lab"), str) and re.fullmatch(r"custom-[0-9a-f]{12}", row.get("lab") or "") else None
        source = row.get("source") if isinstance(row.get("source"), str) and re.match(r"https?://", row.get("source") or "") else ""
        result.append(dict(title=clean_title(row["title"]), url=url[:500], links=links, source=source[:500], difficulty=difficulty_of(row.get("difficulty")) or "",
                           topic=clean_title(row.get("topic") or "")[:80], match=chosen, fit=fit, lab=lab_id))
    return result


# ----------------------------- the learner's own lab -----------------------------
QUOTE_PAIRS = {'"': '"', "'": "'", "“": "”", "‘": "’"}


def masked(text):
    """text with the inside of every quoted string blanked out, same length: notation rules (arrows, missing commas,
    JSON words) are matched on this, so they never reach into a value's own text. An unclosed quote runs to the end."""
    out, i, n = [], 0, len(text)
    while i < n:
        close = QUOTE_PAIRS.get(text[i])
        if close is None:
            out.append(text[i])
            i += 1
            continue
        j = i + 1
        while j < n and text[j] != close:
            j += 2 if text[j] == "\\" else 1
        j = min(j, n)
        out.append(text[i] + "\0" * (j - i - 1) + text[j:j + 1])
        i = j + 1
    return "".join(out)


def outside_quotes(pattern, repl, text):
    """pattern.sub(repl, text), but only where the text isn't inside a quoted string."""
    out, at = [], 0
    for found in pattern.finditer(masked(text)):
        out += [text[at:found.start()], repl(found) if callable(repl) else found.expand(repl)]
        at = found.end()
    return "".join(out) + text[at:]


JSON_WORDS = re.compile(r"\b(?:true|false|null)\b")


def literal(text):
    """A plain value written as JSON or Python: [2, 7], "abc", 'abc', True, null, (1, 2)."""
    text = text.strip()
    if not text:
        raise ValueError("Write a value, e.g. [2, 7, 11] or \"abc\" or 9.")
    try:
        value = json.loads(text)
    except (json.JSONDecodeError, RecursionError):
        try:  # JSON's words as Python's, only where they are words of the notation: 'is true' stays text.
            value = ast.literal_eval(outside_quotes(JSON_WORDS, lambda m: {"true": "True", "false": "False", "null": "None"}[m.group(0)], text))
        except (ValueError, TypeError, SyntaxError, MemoryError, RecursionError):  # TypeError: {{1}}, an unhashable set member.
            raise ValueError(f"“{text[:40]}” isn't a value. Use a number, \"text\", a [list], True/False or None.")
    return plain(value)


def plain(value, depth=0):
    """Plain data the tracer can show: numbers, text, booleans, None, lists and string-keyed dicts."""
    if depth > 3:
        raise ValueError("Values may nest at most three levels deep.")
    if isinstance(value, tuple):
        value = list(value)
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, int):
        if abs(value) > 10 ** 9:
            raise ValueError("Use numbers between -1,000,000,000 and 1,000,000,000.")
        return value
    if isinstance(value, float):
        if value != value or abs(value) > 1e9:
            raise ValueError("Use finite numbers up to 1,000,000,000.")
        return value
    if isinstance(value, str):
        if len(value) > 200:
            raise ValueError("Use at most 200 characters per string.")
        return value
    if isinstance(value, list):
        if len(value) > 200:
            raise ValueError("Use at most 200 items per list.")
        return [plain(v, depth + 1) for v in value]
    if isinstance(value, dict):
        if len(value) > 100 or any(not isinstance(k, str) for k in value):
            raise ValueError("Dictionaries may have at most 100 string keys.")
        return {k: plain(v, depth + 1) for k, v in value.items()}
    raise ValueError("Use numbers, text, True/False, None, lists or dictionaries.")


def split_top(text, sep=","):
    """Split on separators outside brackets and quotes."""
    parts, depth, quote, start = [], 0, None, 0
    for i, ch in enumerate(text):
        if quote:
            if ch == quote and text[i - 1] != "\\":
                quote = None
        elif ch in "\"'":
            quote = ch
        elif ch in "([{":
            depth += 1
        elif ch in ")]}":
            depth -= 1
        elif ch == sep and depth == 0:
            parts.append(text[start:i])
            start = i + 1
    parts.append(text[start:])
    return parts


QUOTES = str.maketrans({"“": '"', "”": '"', "„": '"', "‟": '"', "″": '"', "‘": "'", "’": "'", "‚": "'", "‛": "'", "′": "'"})
# A value in a chain: a number or a word; a "-" is part of it only when it isn't the start of an arrow ("1->-2").
# A chain starts only where a value starts, and no part is ever given back, so "1-1-1-…" is read in one pass.
CHAIN_TOKEN, CHAIN_ARROW = r"(?:-(?!>)|[\w.])++", r"\s*+<?->\s*+"
CHAIN_BODY = (rf"(?<![\w.\-]){CHAIN_TOKEN}(?:{CHAIN_ARROW}{CHAIN_TOKEN})++"
              rf"|(?<![\w.\-])[A-Za-z_]\w*+{CHAIN_ARROW}(?=[\],]|$)")
# "[head -> 1 -> 2]" is one list: its brackets belong to the chain.
CHAIN = re.compile(r"\[\s*+(?:" + CHAIN_BODY + r")\s*+\]|" + CHAIN_BODY)
NULL_WORDS = {"null", "none", "nullptr"}
KINDS = {"linkedlist", "dll", "tree", "linkedlists", "cycle"}
# How a lab compares: the returned node's value, printed lines without trailing spaces, or the first input as solve
# leaves it (a problem that changes its input in place and returns nothing).
ANSWERS = {"node-value", "lines", "in-place"}
SOURCES = {"page", "leetcode"}  # Where a case's expected output came from; without one, the learner wrote it.


def chain_tokens(raw):
    return [t for t in re.split(r"\s*+<?->\s*+", raw.strip().strip("[]").strip()) if t]


def is_number(token):
    return bool(re.fullmatch(r"-?\d+(?:\.\d+)?", token))


def chain_values(tokens):
    """A chain's values without its end marker: NULL / null / None / nullptr, or an x or X after numbers ("1 -> 2 -> X").
    After letters an x is a value: "a -> b -> x" ends with "x"."""
    tokens = list(tokens)
    while tokens and tokens[-1].lower() in NULL_WORDS:
        tokens.pop()
    if len(tokens) > 1 and tokens[-1] in ("x", "X") and all(is_number(t) for t in tokens[:-1]):
        tokens.pop()
    return tokens


def chain_name(part):
    """The parameter a chain names itself ("head -> 1 -> 2"): its first word, when that word is a list's name
    (head, list1, l1…) or the values after it are numbers. In "a -> e -> b" every word is a value."""
    if not re.match(r"[A-Za-z_]\w*\s*<?->", part) or not CHAIN.fullmatch(masked(part.strip())):
        return None
    first, *rest = chain_tokens(part)
    if first.lower() in NULL_WORDS or first.lower() in ("true", "false"):
        return None
    return first if first.lower() in LIST_NAMES or all(is_number(t) for t in chain_values(rest)) else None


def chains(text, labels=()):
    """Linked lists written with arrows ("head -> 1 -> 2 -> null", "1 <-> 2") → JSON lists, and the kind seen.
    Arrows are read only outside quoted strings: "0->2" in quotes is text. A first word is dropped as the list's
    label only when it is one of labels (the name the chain gives itself, or a list's name such as head)."""
    kind = None
    labels = {label.lower() for label in labels}

    def convert(match):
        nonlocal kind
        raw = match.group(0)
        tokens = chain_tokens(raw)
        if tokens and tokens[0].lower() in labels:
            tokens = tokens[1:]  # The label: "head", "list1".
        values = []
        for token in chain_values(tokens):
            try:
                values.append(literal(token))
            except ValueError:
                values.append(token)
        kind = "dll" if "<->" in raw else kind or "linkedlist"
        return json.dumps(values)

    return outside_quotes(CHAIN, convert, text), kind


TEXT_NAMES = {"s", "str", "text", "word", "string", "t", "pattern", "expression", "expr", "exp"}


MISSING_COMMA = re.compile(r"(?<=[\w\]\)\"'])\.?[ \t]+(?=[A-Za-z_]\w*\s*=(?!=))")
# Each call is a name, its parenthesised arguments and at most one comma: one way to read any text, so a near
# miss fails in one pass instead of trying every split of the spaces between calls.
DESIGN_CALLS = re.compile(r"\[\s*((?:[A-Za-z_]\w*\s*\([^()]*\)\s*(?:,\s*)?)+)\]")


def assignments(given, readings=None):
    """An example's input → (parameter names, values, kinds). Judges' notations are read as written:
    "nums = [2, 7], target = 9"; GfG's "arr[] = {1, 2}"; linked lists as "head -> 1 -> 2"; and design
    problems as an operations list with its arguments ("operations = [...]" / "nums = [...]" or two bare lists).
    Each reading beyond the page's own notation is added to readings, to be said in the notes."""
    readings = [] if readings is None else readings
    given = straighten(given)  # A quote that delimits a value becomes plain; one inside a value stays as written.
    # "M = 2 edge = [...]" or "[...]. target = 1": a missing comma between named inputs, never inside a quoted value.
    separated = outside_quotes(MISSING_COMMA, ", ", given)
    if separated != given:
        readings.append("The page separates two named inputs without a comma (by a space or a full stop), so they were read as separate inputs.")
    given = separated
    lines = [line.strip() for line in given.split("\n") if line.strip()]
    balanced = lambda t: (lambda m: sum(m.count(c) for c in "([{") == sum(m.count(c) for c in ")]}"))(masked(t))
    if not all(balanced(line) for line in lines):
        lines = [" ".join(lines)]
    parts = [part.strip() for line in lines for part in split_top(line) if part.strip()]
    named = []
    for part in parts:
        found = re.match(r"^([A-Za-z_]\w*)\s*(?:\[\s*\]\s*)?=\s*(.*)$", part, re.S)
        label = None if found else chain_name(part)
        # After "name =" every word of a chain is a value; a chain that names itself drops that name.
        named.append((found.group(1), found.group(2), ()) if found else (label, part, (label,)) if label else None)
    names, values, kinds = [], [], {}

    def read(name, raw, labels):
        text, kind = chains(raw, labels)
        try:
            value = literal(braced(text))
        except ValueError:
            # "s = (*))": a text input written without quotes, one token long.
            if name.lower() in TEXT_NAMES and re.fullmatch(r"[^\s,\[\]{}\"']+", text.strip()):
                readings.append(f"{name} is written without quotes, so it was read as the text {text.strip()}.")
                return text.strip()
            raise
        if kind == "linkedlist" and raw.strip().startswith("[") and isinstance(value, list) and all(isinstance(v, list) for v in value):
            kind = "linkedlists"  # A list of linked lists, as in "merge k sorted lists".
        if kind:
            kinds[name] = kind
        return value

    calls = DESIGN_CALLS.fullmatch(" ".join(lines))
    if calls:  # A design problem written as calls: [MedianFinder(), addNum(1), findMedian()] → operations and their arguments.
        found = re.findall(r"([A-Za-z_]\w*)\s*\(([^()]*)\)", calls.group(1))
        return ["operations", "arguments"], [[name for name, _ in found], [literal(f"[{args}]") for _, args in found]], {}
    if named and all(named):
        for name, raw, labels in named:
            names.append(name)
            values.append(read(name, raw, labels))
    elif len(lines) == 2 and all(line.startswith("[") for line in lines):
        names, values = ["operations", "arguments"], [literal(lines[0]), literal(lines[1])]
    else:
        names, values = ["value"], [read("value", " ".join(lines), LIST_NAMES)]
    return names, values, kinds


def design_entry(names, values):
    """A design problem's class name, when an example is a list of operations and their arguments."""
    if (len(values) == 2 and isinstance(values[0], list) and values[0] and all(isinstance(x, str) for x in values[0])
            and re.fullmatch(r"[A-Z]\w*", values[0][0]) and isinstance(values[1], list) and len(values[1]) == len(values[0])):
        return values[0][0]
    return None


def signatures(statement, operations):
    """Each operation's parameter names, as the statement writes them: "void push(int x)" → ["x"]."""
    found = {}
    for op in dict.fromkeys(operations):
        match = re.search(r"\b" + re.escape(op) + r"\s*\(([^)]*)\)", statement)
        if match:
            names = [re.split(r"\s+", a.strip())[-1].strip("&*[]") for a in match.group(1).split(",") if a.strip()]
            found[op] = [n for n in names if re.fullmatch(r"[A-Za-z_]\w*", n) and not keyword.iskeyword(n)]
    return found


LIST_NAMES = {"head", "head1", "head2", "heada", "headb", "l1", "l2", "list1", "list2", "linkedlist", "ll", "lst"}
TREE_NAMES = {"root", "root1", "root2", "tree", "t1", "t2"}
KIND_WORDS = {"linkedlist": "a linked list", "dll": "a doubly linked list", "tree": "a binary tree", "linkedlists": "a list of linked lists"}


def guess_kinds(params, statement, sample, found=None):
    """Which parameters are linked lists or trees: as written with arrows, or named so in a problem about them."""
    kinds, text = dict(found or {}), statement.lower()
    for i, name in enumerate(params):
        value = sample[i] if i < len(sample) else None
        if name in kinds or not isinstance(value, list):
            continue
        key = name.lower().replace("_", "")
        if key in LIST_NAMES and "linked list" in text:
            kinds[name] = "dll" if "doubly" in text else "linkedlist"
        elif key in TREE_NAMES and ("tree" in text or "bst" in text):
            kinds[name] = "tree"
        elif key in ("heads", "lists") and "linked list" in text and all(isinstance(v, list) for v in value):
            kinds[name] = "linkedlists"
    for i, name in enumerate(params):
        value = sample[i] if i < len(sample) else None
        if name.lower() == "pos" and type(value) is int and ("cycle" in text or "loop" in text) and any(k in ("linkedlist", "dll") for k in kinds.values()):
            kinds[name] = "cycle"
    return kinds


def braced(value):
    """{1, 2, 3} in a C-style example is a list, not a set (a dictionary keeps its braces)."""
    value = value.strip()
    if value.startswith("{") and value.endswith("}") and ":" not in value:
        return "[" + value[1:-1] + "]"
    return value


EXAMPLE_LABEL = re.compile(r"(input|output)\s*[:：]", re.I)
EXPLANATION_WORD, COLON_NEXT, SPACES = re.compile(r"explanation", re.I), re.compile(r"\s*:"), re.compile(r"\s*")


def example_pairs(text):
    """The (input, output) pairs of pasted examples: an input runs from "Input:" to the next "Output:", and an
    output to the end of its line or to an "Explanation:" on that line. Read label by label in one pass."""
    labels, pairs, after = list(EXAMPLE_LABEL.finditer(text)), [], 0
    for i, start in enumerate(labels):
        if start.group(1).lower() != "input" or start.start() < after:
            continue
        # An input holds at least one character after its own spaces; only when no later "Output:" allows that is
        # an "Output:" straight after those spaces taken (with a blank input), as the single pattern did.
        filled, chosen, fallback = SPACES.match(text, start.end()).end(), None, None
        for label in (labels[k] for k in range(i + 1, len(labels))):
            if label.group(1).lower() != "output" or label.end() == len(text):
                continue  # Nothing at all after an output label: the input runs on to a later one.
            if label.start() > filled:
                chosen = label
                break
            if label.start() == filled > start.end():
                fallback = label
        label = chosen or fallback
        if not label:
            break  # No output follows this input, so none follows a later one either.
        begin = min(SPACES.match(text, label.end()).end(), len(text) - 1)  # Only spaces after it: a blank output, which fails as a value.
        end = text.find("\n", begin + 1)
        end = len(text) if end < 0 else end
        for word in EXPLANATION_WORD.finditer(text, begin + 1, end):
            gap = word.start()
            while gap > begin + 1 and text[gap - 1].isspace():
                gap -= 1
            if gap < word.start() and COLON_NEXT.match(text, word.end()):
                end = gap
                break
        pairs.append((text[start.end():label.start()].strip(), text[begin:end]))
        after = end
    return pairs


def parse_examples(text):
    """Examples as judges print them ("Input: nums = [2,7], target = 9 / Output: [0,1]") → parameter
    names and cases. Only what the text states is read; nothing is completed or guessed, and any reading
    beyond the notation (a missing comma, unquoted text) is said in notes."""
    if not isinstance(text, str) or len(text) > 20000:
        raise ValueError("Paste up to 20,000 characters of examples.")
    blocks = example_pairs(text)
    if not blocks:
        raise ValueError("No “Input: … Output: …” pairs were found. Paste examples in that form, or add cases by hand.")
    params, cases, kinds, notes = None, [], {}, []
    for number, (given, output) in enumerate(blocks[:MAX_CASES], 1):
        said = []
        names, values, found = assignments(given, said)
        kinds.update(found)
        if params is None:
            params = names
        elif names != params:
            raise ValueError(f"Example {number} names different parameters ({', '.join(names)}) than example 1 ({', '.join(params)}).")
        cases.append(dict(name=f"Example {number}", args=values, expected=literal(chains(straighten(" ".join(output.split())), LIST_NAMES)[0])))
        if said:
            notes.append(f"Example {number}: {' '.join(said)}")
    entry = design_entry(params, cases[0]["args"]) if cases else None
    return dict(params=params, cases=cases, kinds=kinds, entry=entry or "solve", **({"notes": notes} if notes else {}))


INVISIBLE = dict.fromkeys(map(ord, "​⁠﻿­"))  # Zero-width spaces, word joiners, byte-order marks, soft hyphens.


class PageBlocks(html.parser.HTMLParser):
    """A page as text blocks (headings, paragraphs, list items, table rows, code blocks, images), skipping
    scripts, controls, forms, navigation and asides. Code blocks keep their exact whitespace. End tags a page
    leaves out are implied as HTML5 implies them, so the text of an unclosed <li> or <p> is kept, in order."""
    BLOCKS = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "dt", "dd", "blockquote", "tr"}
    SKIP = {"script", "style", "noscript", "svg", "button", "nav", "footer", "aside", "select", "option", "label", "textarea", "form", "template", "iframe", "math"}
    LISTS = {"ul", "ol", "dl", "menu"}
    ITEMS = {"li": {"li"}, "dt": {"dt", "dd"}, "dd": {"dt", "dd"}}  # The open item a new one closes.

    def __init__(self):
        super().__init__(convert_charrefs=True)
        # Open blocks are (tag, text parts, lists open when it began); counts make "is a <tag> open?" one lookup.
        self.blocks, self._open, self._skip, self._marks, self._count, self._lists = [], [], 0, [], {}, 0
        self.raw = {}  # Index in blocks → that block's text before its spaces were collapsed (not for code blocks).
        self._resume = False  # A paragraph a list or block interrupted, whose text goes on after it, until </p>.

    def add(self, text):
        if self._skip:
            return
        if not self._open and self._resume and text.strip():
            self.open("p")  # "<p>a<ul>…</ul>b</p>": a browser shows b after the list, so it is kept, as its own paragraph.
            self._resume = False
        if self._open:
            self._open[-1][1].append(text)

    def open(self, tag):
        self._open.append((tag, [], self._lists))
        self._count[tag] = self._count.get(tag, 0) + 1

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif self._skip:
            return
        elif tag == "br":
            self.add("\n")
        elif tag in ("sup", "sub") and self._open:
            # 10<sup>5</sup> is 10^5, not 105; arr<sub>2</sub> is arr_2, not arr2. An empty one marks nothing.
            self.add("^" if tag == "sup" else "_")
            self._marks.append((self._open[-1][1], len(self._open[-1][1]) - 1))
        elif tag in ("td", "th") and self._open and self._open[-1][0] == "tr" and self._open[-1][1]:
            self.add(" | ")
        elif tag == "img":
            self.blocks.append(("img", dict(attrs).get("alt") or ""))
        elif tag in self.LISTS:
            if self._open and self._open[-1][0] == "p":
                self.close_to(len(self._open) - 1)  # A list ends an open paragraph; its text may go on after the list.
                self._resume = not self._open
            self._lists += 1
        elif tag in self.BLOCKS:
            closes = self.ITEMS.get(tag, ())
            for i in range(len(self._open) - 1, -1, -1) if closes else ():  # An open item of the same list ends here...
                name, _, lists = self._open[i]
                if lists < self._lists or (name not in closes and name != "p"):
                    break
                if name in closes:
                    self.close_to(i)
                    break
            if self._open and self._open[-1][0] == "p":
                self.close_to(len(self._open) - 1)  # ...and so does an open paragraph.
                self._resume = tag != "p" and not self._open
            self.open(tag)

    def handle_startendtag(self, tag, attrs):
        if tag in ("br", "img"):
            self.handle_starttag(tag, attrs)

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self._skip = max(0, self._skip - 1)
            return
        if self._skip:
            return
        if tag == "p":
            self._resume = False  # The paragraph's own end: nothing after it continues it.
        if tag in ("sup", "sub") and self._marks:
            parts, at = self._marks.pop()
            if not "".join(parts[at + 1:]).strip():
                parts[at] = ""
        elif tag in self.LISTS and self._lists:
            i = len(self._open)
            while i and self._open[i - 1][2] >= self._lists:
                i -= 1
            self.close_to(i)  # A list's end closes whatever is still open inside it.
            self._lists -= 1
        elif tag in self.BLOCKS and self._count.get(tag):
            i = len(self._open) - 1
            while self._open[i][0] != tag:
                i -= 1
            self.close_to(i)

    def close_to(self, index):
        """Close the open blocks from the innermost out to index, keeping each one's text."""
        while len(self._open) > index:
            name, parts, _ = self._open.pop()
            self._count[name] -= 1
            raw = "".join(parts)
            if name == "pre" or self._count.get("pre"):
                text = "\n".join(line.rstrip() for line in raw.strip("\n").split("\n"))  # Code keeps its indentation.
            else:
                text = "\n".join(" ".join(line.split()) for line in raw.split("\n")).strip()
            if text.strip():
                if name != "pre" and not self._count.get("pre"):
                    self.raw[len(self.blocks)] = raw
                self.blocks.append((name, text))

    def handle_data(self, data):
        self.add(data.translate(INVISIBLE))

    def close(self):
        super().close()
        self.close_to(0)  # The end of the page ends every block still open.


EXAMPLE_HEAD = re.compile(r"^example(?:\s*\d+)?\s*:?$", re.I)
HEADINGS = {"h2", "h3", "h4", "h5", "h6"}
# Sections of a problem worth keeping, with their own heading: "Note: …", "Your Task:", "Input Format", …
SECTION_LABELS = re.compile(r"^(input format|input description|input|output format|output description|output|notes?|your task|task|"
                            r"expected (?:time|auxiliary|space)[\w ]*|follow[- ]?up|function description|parameters|return value|returns|"
                            r"edge cases?|special cases?|assumptions?|clarifications?)\s*[:：]\s*(.*)$|^(input format|output format|input description|output description|notes?|your task|follow[- ]?up)\s*$", re.I | re.S)
# Page furniture and spoilers: they end whatever section they interrupt and are never imported.
STOP_LABELS = re.compile(r"^(hints?|editorial|solutions?|approach(es)?|discuss\w*|comments?|similar (problems|questions)|related (problems|articles|topics)|companies|company tags|"
                         r"topics?|tags|now your turn|still unsure|let.s go through|explain (the )?problem|video|submissions?|accuracy|recommended|complexity analysis|"
                         r"code|implementation|dry run|(naive|brute[- ]?force|better|optimal|efficient|expected) (approach|solution)|intuition|algorithm|how to solve|"
                         r"try (it|more)|practice (this|now)|login|sign (in|up)|upgrade|subscribe|share|table of contents?|in this article|on this page)\b", re.I)
# The prompts among them that a page writes as a sentence ("Still unsure what the problem is asking?"). The rest end
# the problem only as a heading or a bare label, so "Implementation must run in O(n)." stays in the statement.
STOP_PROMPTS = re.compile(r"^(now your turn|still unsure|let.s go through|explain (the )?problem|table of contents?|in this article|on this page|try (it|more) (yourself|on)|practice (this|now))\b", re.I)
SAMPLE_INPUT = re.compile(r"^sample\s+input\s*#?\s*\d*\s*:?", re.I)
SAMPLE_OUTPUT = re.compile(r"^sample\s+output\s*#?\s*\d*\s*:?", re.I)
# Page furniture as whole short lines ("Login to save progress", "You haven't submitted yet"), never a problem's own
# words: "each cookie j has a size" or "users subscribe to topics" are kept.
JUNK_LINE = re.compile(r"^(?:log ?in|sign (?:in|up)|subscribe|upgrade)(?: (?:now|here|to [\w ]{1,40}|for [\w ]{1,30}))?[.!]?$"
                       r"|\b(?:we use cookies|accept (?:all )?cookies|cookie (?:policy|settings|preferences)|all rights reserved|download (?:the|our) app|"
                       r"you haven.t submitted|submit your code|run against hidden test cases)\b|©", re.I)
# A heading that introduces the statement itself ("Description", "Problem Statement") rather than ending it.
INTRO_HEAD = re.compile(r"^(?:problem(?:\s+(?:statement|description))?|description|statement|question)\s*[:：]?$", re.I)
MAX_STATEMENT = 12000


SINGLE_OPEN, SINGLE_CLOSE = re.compile(r"(?<=[=\[,(\s])[‘‚‛]"), re.compile(r"’(?=\s*(?:[,\])]|$))")


def straighten(text):
    """Typographic quotes that delimit a value become plain quotes; quotes inside a value are left as written.
    A ‘ opens a value after =, [, ( , or a space, and the value ends at the first ’ on its line that a , ] ) or the end
    follows. Each mark is looked at once, so a line of unmatched ‘ costs one pass."""
    text = re.sub(r"[“„‟″]([^“”„‟″]*)[”″]", r'"\1"', text)
    out, at, close, line_end = [], 0, None, -1
    for opener in SINGLE_OPEN.finditer(text):
        if opener.start() < at:
            continue
        if close is None or close.start() < opener.end():
            close = SINGLE_CLOSE.search(text, opener.end())  # The first ’ that can end a value, from here on.
            if not close:
                break
        if line_end < opener.end():
            line_end = text.find("\n", opener.end())
            line_end = len(text) if line_end < 0 else line_end
        if close.start() > line_end:
            continue  # Its line has no closing ’; a later line's opener may still use that one.
        out += [text[at:opener.start()], "'", text[opener.end():close.start()], "'"]
        at = close.end()
    return "".join(out) + text[at:]


OUT_LABEL = r"output\s*(?:\([^)\n]{0,80}\)\s*)?[:：=]"  # "Output:", "Output =", "Output(value at returned node):"


def split_example(lines):
    """One example's text → its raw input, output, explanation, and any note in the output's label
    ("Output(value at returned node): 7"), exactly as written."""
    text = "\n".join(lines)
    out = OUT_LABEL if re.search(OUT_LABEL, text, re.I) else r"(?<![\w-])result\s*[:：]"  # Some pages label it "Result:".
    # Each part ends where the next label begins; the spaces before that label are trimmed by clean() below, so
    # they needn't be matched (matching them made a run of blank lines quadratic).
    given = re.search(rf"input\s*[:：]\s*(.*?)(?={out}|\Z)", text, re.I | re.S)
    shown = re.search(rf"{out}[ \t]*(.*?)(?=explanation\s*[:：]|\Z)", text, re.I | re.S)
    label = re.search(r"output\s*\(([^)\n]{0,80})\)\s*[:：=]", text, re.I)
    why = re.search(r"explanation\s*[:：]\s*(.*)\Z", text, re.I | re.S)
    clean = lambda m: "\n".join(line.rstrip() for line in m.group(1).strip().split("\n")) if m else ""
    given, output, why = clean(given), clean(shown), clean(why)
    # An output ends at a later labelled line ("Expalantion: …", as one page spells it): that line explains, it isn't output.
    rows = output.split("\n")
    cut = next((i for i, line in enumerate(rows[1:], 1) if rows[0].strip() and re.match(r"\s*[A-Za-z][\w ]{0,20}[:：]\s", line)), None)
    if cut is not None:
        trailing = re.sub(r"^ex\w{0,3}la\w{0,3}tion\s*[:：]\s*", "", "\n".join(rows[cut:]).strip(), flags=re.I)
        output, why = "\n".join(rows[:cut]).rstrip(), "\n".join(x for x in [trailing, why] if x)
    return given, output, why, label.group(1).strip() if label else ""


NUMBERS = re.compile(r"-?\d+(?:\.\d+)?(?:[ \t]+-?\d+(?:\.\d+)?)+")


def read_output(output, params, readings=None):
    """An output as the page writes it → (value, how): how is None for a plain value, or the notation read —
    "spaced" (numbers separated by spaces, as judges print a returned list), "word" (a bare word: text),
    "assigned" (one value written as name = value) or "named" (one value per input, named after them).
    Anything else is not read. Readings inside an assignment (a missing comma…) are added to readings."""
    text = output.strip()
    if "\n" not in text and re.match(r"[A-Za-z_]\w*\s*=(?!=)", text):
        said = []
        try:
            names, values, _ = assignments(text, said)
        except ValueError:
            names = None
        if names and (len(names) == 1 or names == params) and readings is not None:
            readings.extend(said)
        if names and len(names) == 1:
            return values[0], "assigned"  # "head = [1, 2]": the page names what is returned.
        if names and len(names) > 1 and names == params:
            return values, "named"
    single = chains(straighten(text), LIST_NAMES)[0] if "\n" not in text else ""
    for attempt in [single] + [text, text.translate(QUOTES)]:
        try:
            value = literal(attempt)
        except ValueError:
            continue
        if attempt is not single and attempt != text and readings is not None:
            readings.append("The output's typographic quotes (“ ” ‘ ’) were read as plain quotes.")
        return value, None
    if NUMBERS.fullmatch(text):
        return [literal(x) for x in text.split()], "spaced"
    if re.fullmatch(r"[A-Za-z][A-Za-z0-9_'-]*", text):
        return text, "word"
    raise ValueError("not a value")


def example_case(example, number):
    """One example → its record (raw text and status) and, when the values can be read, a case.
    The input is read as written; an output that isn't a value (a printed pattern, prose) is left for the
    learner, never guessed."""
    name = example.get("name") or f"Example {number}"
    if example.get("sample"):
        given, output, why, label = "\n".join(example["input"]), "\n".join(example["output"]), "", ""
    else:
        given, output, why, label = split_example(example["lines"])
    record = dict(name=name, input=given, output=output, explanation=why, status="parsed", issue=None, images=example.get("images", 0))
    if not given:
        record.update(status="unparsed", issue="The example has no input to read.")
        return record, None
    # As written; then with every typographic quote plain; then, for a tree written in level order, with each
    # bare N (the page's missing node) read as null. Each reading beyond the page's notation is said on the example.
    nulls = lambda text: outside_quotes(re.compile(r"(?<=[\[,])(\s*)N(?=\s*[,\]])"), r"\1null", text)
    attempts = [(straighten(given), False), (given.translate(QUOTES), False)] + ([(nulls(given), True)] if nulls(given) != given else [])
    for attempt, (text, as_null) in enumerate(attempts):
        said = []
        try:
            names, values, kinds = assignments(text, said)
        except ValueError:
            continue
        if as_null and not any(name.lower() in TREE_NAMES for name in names):
            continue  # Only a tree's level order writes a missing node as N.
        if as_null:
            output = nulls(output)
            said.append("Each N in this example was read as null: the page writes a missing tree node as N.")
        elif attempt == 1:
            said.insert(0, "Its typographic quotes (“ ” ‘ ’) were read as plain quotes.")
        for reading in said:
            say(record, reading)
        break
    else:
        reason = "it is written as standard input lines" if example.get("sample") or "\n" in given.strip() else "it isn't written as values"
        record.update(status="unparsed", issue=f"The input couldn't be read as values ({reason}). Enter this case by hand.")
        return record, None
    expected, missing = None, False
    if re.search(r"node", label, re.I) and re.search(r"value", label, re.I):
        record["node_value"] = True  # "Output(value at returned node): 7": the page shows the returned node's value.
    if not output:
        missing = True
        record["status"] = "partial"
        say(record, "The page shows this output as an image, which can't be imported. Open the original to see it and write the expected output yourself."
            if example.get("images") else "The page shows no output value. Write the expected output yourself.")
    else:
        said = []
        try:
            expected, how = read_output(output, names, said)
        except ValueError:
            missing = True
            record["status"] = "partial"
            say(record, "The output isn't a single value (for example a printed pattern). Write the expected output yourself.")
        else:
            for reading in said:
                say(record, reading)
            if how:
                record["read"] = how
                say(record, {"spaced": "The output is written as numbers separated by spaces, as judges print a returned list, so it was read as a list.",
                             "word": "The output is a bare word, so it was read as text.",
                             "assigned": "The output is written as an assignment (name = value), so the value after = was read as the expected output.",
                             "named": f"The output names a value for each input ({', '.join(names)}), so it was read as a list in that order."}[how])
    case = dict(name=name, params=names, args=values, expected=expected, missing=missing, kinds=kinds, explanation=why[:1000])
    return record, case


def say(record, text):
    """Add a sentence to what an example's record says about how it was read."""
    record["issue"] = f"{record['issue']} {text}" if record["issue"] else text


STATEMENT_KEYS = {"problem question", "problemquestion", "problem statement", "problemstatement", "statement", "description", "content", "question", "body", "question content", "questioncontent"}


def embedded_problem(page):
    """A problem a page keeps in its embedded data rather than its markup (GeeksforGeeks' Next.js page data:
    problem_name plus problem_question as HTML), rebuilt as a small page to read the same way."""
    best = None

    def visit(node, depth=0):
        nonlocal best
        if depth > 60 or best:
            return
        if isinstance(node, list):
            for item in node:
                visit(item, depth + 1)
        elif isinstance(node, dict):
            title = next((v for k, v in node.items() if normal(k) in TITLE_KEYS and isinstance(v, str) and 2 <= len(v.strip()) <= 160), "")
            body = next((v for k, v in node.items() if normal(k) in STATEMENT_KEYS and isinstance(v, str) and len(v) > 60 and re.search(r"<(p|div|pre|span|strong)\b", v)), "")
            if title and body:
                best = f"<h1>{html.escape(title)}</h1>{body}"
                return
            for value in node.values():
                visit(value, depth + 1)

    for value in embedded_json(page):
        visit(value)
    return best


BOUNDS = [  # GeeksforGeeks' constraint fields: (low key, high key, what they bound).
    ("min_size", "max_size", "{name}.size()"), ("element_min_value", "element_max_value", "{name}[i]"),
    ("element_min_length", "element_max_length", "{name}[i].length()"), ("min_value", "max_value", "{name}"),
    ("min_length", "max_length", "{name}.length()"), ("min_nodes", "max_nodes", "number of nodes"),
    ("node_min_value", "node_max_value", "node value"), ("min_rows", "max_rows", "number of rows"), ("min_cols", "max_cols", "number of columns"),
]
BOUND_META = {"show_cross_args_in_preview", "node_value_type"}


def data_constraints(page):
    """Constraints a page keeps as data and draws in the browser (GeeksforGeeks' input_format), written out
    bound by bound exactly as stored. Fields that aren't understood are left out and counted, never guessed."""
    found = []

    def visit(node, depth=0):
        if depth > 60 or found:
            return
        if isinstance(node, dict):
            form = node.get("input_format")
            if isinstance(form, dict) and isinstance(form.get("constraints"), list) and form.get("show_constraints_in_problem"):
                found.append(form)
                return
            for value in node.values():
                visit(value, depth + 1)
        elif isinstance(node, list):
            for value in node:
                visit(value, depth + 1)

    for value in embedded_json(page):
        visit(value)
    if not found:
        return [], 0
    names = [re.match(r"\s*([A-Za-z_]\w*)", part) for part in str(found[0].get("arguments") or "").split("&!//!&")]
    lines, skipped = [], 0
    for i, rule in enumerate(found[0]["constraints"]):
        name = names[i].group(1) if i < len(names) and names[i] else None
        if not isinstance(rule, dict) or not name:
            skipped += 1
            continue
        known = set(BOUND_META)
        # "k ≤ arr.size()": a value bounded by another argument, completing a missing upper bound.
        cross = [c for c in rule.get("cross_argument_constraints") or [] if isinstance(c, dict) and c.get("sourceProperty") == "value"
                 and c.get("operator") == "≤" and c.get("targetArg") and c.get("targetProperty") in ("size", "value")]
        if len(cross) == 1 == len(rule["cross_argument_constraints"]) and rule.get("max_value") in (None, ""):
            known.add("cross_argument_constraints")
        for low, high, what in BOUNDS:
            lo, hi = rule.get(low), rule.get(high)
            if what == "{name}" and "cross_argument_constraints" in known:
                hi = cross[0]["targetArg"] + (".size()" if cross[0]["targetProperty"] == "size" else "")
            if lo in (None, "") and hi in (None, ""):
                continue
            known |= {low, high}
            subject = what.format(name=name)
            lines.append(f"{lo} ≤ {subject} ≤ {hi}" if lo not in (None, "") and hi not in (None, "") else f"{subject} ≥ {lo}" if lo not in (None, "") else f"{subject} ≤ {hi}")
        skipped += len(set(rule) - known)
    return lines, skipped


def problem_from_page(page, url=""):
    """A problem page → its title, statement, constraints and example cases, as the page states them.
    The statement runs from the page's main heading to its first example; any section other than
    examples and constraints (hints, quizzes, editorials, comments) is not read."""
    try:
        return problem_from_blocks(page, url)
    except ValueError:
        stored = embedded_problem(page)
        if not stored:
            raise
        return problem_from_blocks(stored, url, data_constraints(page))


def problem_from_blocks(page, url, stored_constraints=([], 0)):
    """The problem between the page's main heading and its first unrelated section: the description, labelled
    sections (input/output format, notes, task, expected complexity, follow-up), constraints and examples, each
    example keeping its own input, output and explanation."""
    reader = PageBlocks()
    reader.feed(page)
    reader.close()
    blocks = reader.blocks
    start = next((i for i, (tag, _) in enumerate(blocks) if tag == "h1"), None)
    if start is None:
        raise ValueError("No problem was found on that page: it has no main heading to start from.")
    title = clean_title(blocks[start][1])
    statement, sections, constraints, examples, images = [], [], [], [], 0
    sources = []  # Each statement line's block as the page wrote it (None for a code block), for a drawn pattern's spacing.
    mode, current = "statement", None
    started = lambda: bool(statement or examples or sections or constraints)
    junk = 0
    for index, (tag, text) in enumerate(blocks[start + 1:], start + 1):
        if mode == "other" and tag not in HEADINGS:
            continue  # Inside a quiz, an editorial or page furniture: only a new titled section can follow.
        if tag == "img":
            if re.search(r"\b(icon|logo|avatar|badge|emoji|spinner)\b", text, re.I):
                continue  # A link's icon or a logo, not a figure of the problem.
            images += 1
            if mode in ("example", "sample-input", "sample-output") and examples:
                examples[-1]["images"] = examples[-1].get("images", 0) + 1
            continue
        if tag != "pre" and len(text) <= 160 and JUNK_LINE.search(text):
            junk += 1  # Sign-in prompts, cookie banners, a judge's "submit your code" panel.
            continue
        first, _, rest = text.partition("\n")
        if tag in HEADINGS and INTRO_HEAD.match(text.strip()) and not examples:
            mode = "statement"  # "Description", "Problem Statement": the statement follows.
            continue
        headingish = tag in HEADINGS or (tag in {"p", "dt", "blockquote"} and len(text) <= 48 and text.rstrip().endswith((":", "：")))
        label_text = text.strip()
        if (STOP_LABELS.match(label_text) and (headingish or len(label_text.split()) <= 2)) or (STOP_PROMPTS.match(label_text) and len(label_text) <= 90):
            if started() or tag in HEADINGS:
                mode = "other"  # Hints, quizzes, editorials, comments and page prompts are not the problem.
            continue
        if (headingish or EXAMPLE_HEAD.match(text)) and re.match(r"examples?\b", text, re.I):
            named = re.match(r"(example\s*\d+)", text, re.I)
            examples.append({"lines": [], "name": named.group(1).capitalize() if named else f"Example {len(examples) + 1}", "headed": True})
            mode = "example"
            continue
        if SAMPLE_INPUT.match(first):
            examples.append({"sample": True, "input": [], "output": [], "name": f"Sample {len(examples) + 1}"})
            mode = "sample-input"
            remainder = SAMPLE_INPUT.sub("", text, count=1).strip("\n")
            if remainder.strip():
                examples[-1]["input"].append(remainder)
            continue
        if SAMPLE_OUTPUT.match(first) and examples and examples[-1].get("sample"):
            mode = "sample-output"
            remainder = SAMPLE_OUTPUT.sub("", text, count=1).strip("\n")
            if remainder.strip():
                examples[-1]["output"].append(remainder)
            continue
        if re.match(r"constraints?\s*(?:[:：]\s*)?$", first, re.I) or (headingish and re.match(r"constraints?\b", text, re.I)):
            mode = "constraints"
            constraints += [c.strip() for c in rest.split("\n") if c.strip()]
            continue
        label = SECTION_LABELS.match(first)
        bare_io = bool(label) and (label.group(1) or "").lower() in ("input", "output")
        if label and (headingish or label.group(2) is not None) and not (bare_io and tag not in HEADINGS):  # "Input: …" lines are example text.
            heading = (label.group(1) or label.group(3) or "").strip()
            body = "\n".join(x for x in [(label.group(2) or "").strip(), rest.strip()] if x)
            current = {"heading": heading[:1].upper() + heading[1:], "text": body}
            sections.append(current)
            mode = "section"
            continue
        if tag in HEADINGS:
            if started():
                mode = "other"  # Any other titled section once the problem has begun.
            continue
        line = ("- " if tag == "li" else "") + text
        if mode in ("statement", "constraints") and re.match(r"input\s*[:：]", text, re.I):
            examples.append({"lines": [text], "name": f"Example {len(examples) + 1}"})  # An example without a heading.
            mode = "example"
        elif mode == "statement":
            statement.append(line)
            sources.append(reader.raw.get(index))
        elif mode == "example":
            if re.match(r"input\s*[:：]", text, re.I) and any(re.match(r"input\s*[:：]", seen, re.I) for seen in examples[-1]["lines"]):
                examples.append({"lines": [], "name": f"Example {len(examples) + 1}"})  # Several examples under one heading.
            examples[-1]["lines"].append(text)
        elif mode in ("sample-input", "sample-output"):
            examples[-1]["input" if mode == "sample-input" else "output"].append(text)
        elif mode == "constraints":
            constraints.append(text.strip())
        elif mode == "section":
            current["text"] = (current["text"] + "\n" + line).strip()

    notes = []
    if junk:
        notes.append(f"Left out {junk} line{'s' if junk > 1 else ''} of page text that aren't part of the problem (sign-in, cookie or submission prompts).")
    # A bare "Input:" with nothing after it (a judge's custom-input box) isn't an example; a titled one is kept.
    examples = [e for e in examples if e.get("headed") or e.get("sample") or e.get("images") or any(split_example(e["lines"])[:3])]
    sections = [section for section in sections if section["text"]]
    constraints = [re.sub(r"^[-•·*▪◦]\s+", "", c) for c in constraints if c]  # A list marker isn't part of a constraint.
    stored, unread = stored_constraints
    if not constraints and (stored or unread):
        constraints = stored  # Drawn by the page from its data; written out bound by bound as stored.
        if stored:
            notes.append("The constraints come from the page's data (the site draws them in the browser), written out bound by bound.")
        if unread:
            notes.append(f"{unread} constraint field{'s' if unread > 1 else ''} in the page's data couldn't be read. Open the original for the full constraints.")
    if not title or (not statement and not examples):
        raise ValueError("No problem statement was found on that page.")

    records, params, cases, found_kinds, seen = [], None, [], {}, set()
    for number, example in enumerate(examples[:MAX_CASES], 1):
        record, case = example_case(example, number)
        if case:
            key = json.dumps([case["args"], case["expected"]])
            if key in seen:
                record.update(status="duplicate", issue="The same input and output as an earlier example, so it wasn't added again.")
                case = None
            elif params is None:
                params = case["params"]
            elif case["params"] == ["value"] and len(params) == 1:
                case["params"] = params  # "2" after "n = 3": the page leaves the one input unnamed.
                say(record, f"The page doesn't name this input, so it was read as {params[0]}, the one input the first example names.")
            elif case["params"] != params:
                record.update(status="unparsed", issue=f"It names different inputs ({', '.join(case['params'])}) than the first example ({', '.join(params)}).")
                case = None
        if case:
            seen.add(json.dumps([case["args"], case["expected"]]))
            cases.append({k: case[k] for k in ("name", "args", "expected", "missing", "explanation")})
            found_kinds.update(case["kinds"])
        records.append(record)
    if len(examples) > MAX_CASES:
        notes.append(f"The page has {len(examples)} examples; the first {MAX_CASES} were read.")
    by_name = {r["name"]: r for r in records}
    # One list contract: once an output is a space-separated list, a lone number in another example is a list of one.
    if any(r.get("read") == "spaced" for r in records):
        for case in cases:
            record = by_name.get(case["name"], {})
            if type(case["expected"]) in (int, float) and NUMBERS.fullmatch(f"{record.get('output', '')} 0".strip()):
                case["expected"] = [case["expected"]]
                record["read"] = "spaced"
                say(record, "The output is a single number in a list of numbers, so it was read as a list of one, like the other examples.")
    # The page never names the input. Only when the statement speaks of one structure (a binary tree, whose level
    # order is its root, or a linked list, whose values are its head) is it named and read as that, and said so.
    lowered, renamed = " ".join(statement).lower(), None
    if params == ["value"]:
        tree, listed = "binary tree" in lowered or "bst" in lowered, "linked list" in lowered
        numbers = cases and isinstance(cases[0]["args"][0], list) and all(x is None or type(x) is int for x in cases[0]["args"][0])
        if numbers and tree != listed:
            renamed = params = ["root" if tree else "head"]
            notes.append(f"The page doesn't name the input. The statement speaks of a {'binary tree' if tree else 'linked list'} (and no "
                         f"{'linked list' if tree else 'tree'}), so it is called {params[0]} here and read as one. That is an inference, not the page's words.")
        elif numbers and tree:
            notes.append("The page doesn't name the input, and the statement speaks of both a linked list and a tree, so which one the input is "
                         "isn't assumed: it stays a plain list called value. Choose its kind if you know it.")
        else:
            notes.append("The page doesn't name the input, so it is called value here.")
    # A printed pattern drawn in text in the statement ("for N = 5, the pattern should look like: …") is a case from the page.
    pattern, unread, drawn_blocks = statement_pattern(statement, params, cases, sources)
    for at, shown in drawn_blocks.items():
        statement[at] = shown  # The pattern as the page draws it, its &nbsp; indentation kept.
    if pattern:
        cases.append(pattern)
        notes.append(f"The statement draws the pattern for {params[0]} = {pattern['args'][0]} in text, so it became a case: the printed lines, as a list of strings.")
    elif unread:
        notes.append(unread)
    # Cases with a known output are the lab; examples whose output can't be read stay listed above, by name.
    unfilled = []
    if any(not case["missing"] for case in cases) and any(case["missing"] for case in cases):
        left = [case["name"] for case in cases if case["missing"]]
        unfilled = [case for case in cases if case["missing"]]
        cases = [case for case in cases if not case["missing"]]
        notes.append(f"{', '.join(left)} {'was' if len(left) == 1 else 'were'} not added as {'a case' if len(left) == 1 else 'cases'}: the page doesn't give {'its' if len(left) == 1 else 'their'} output as a value. Add {'it' if len(left) == 1 else 'them'} by hand if you like.")
    for record in records:
        if record["issue"] and record["status"] != "duplicate":
            notes.append(f"{record['name']}: {record['issue']}")

    # The statement as the page structures it: description, its labelled sections, constraints, and any example
    # that couldn't become a case, word for word, so nothing the page says is lost.
    parts = ["\n".join(statement)]
    parts += [f"{section['heading']}:\n{section['text']}" for section in sections]
    if constraints:
        parts.append("Constraints:\n" + "\n".join("- " + c for c in constraints))
    for record in records:
        if record["status"] in ("partial", "unparsed"):
            drawn = "Output: (shown as an image on the page)" if record["images"] else ""
            shown = [f"Input: {record['input']}" if record["input"] else "", f"Output: {record['output']}" if record["output"] else drawn, f"Explanation: {record['explanation']}" if record["explanation"] else ""]
            parts.append(f"{record['name']}, as written on the page:\n" + "\n".join(x for x in shown if x))
    text = "\n\n".join(part for part in parts if part.strip())
    truncated = len(text) > MAX_STATEMENT
    if truncated:
        text = text[:MAX_STATEMENT].rsplit("\n", 1)[0]
        notes.append(f"The problem text is longer than {MAX_STATEMENT:,} characters; the rest wasn't imported. Open the original for all of it.")
    drawn = sum(r["images"] for r in records if r["images"] and not r["output"])  # Already said on their examples.
    if images - drawn > 0:
        unsaid = images - drawn
        notes.append(f"The problem includes {unsaid}{' more' if drawn else ''} image{'s' if unsaid > 1 else ''} (a diagram or figure) that can't be imported. Open the original to see {'them' if unsaid > 1 else 'it'}.")
    if not cases:
        notes.append("No example could be read as a case. Add cases by hand.")
    elif any(case["missing"] for case in cases) and re.search(r"\bprint", text, re.I):
        notes.append("This problem prints its answer, but a lab checks what solve returns: return the printed lines "
                     "(for example a list of strings) and write that list as each expected output.")
    entry = design_entry(params or [], cases[0]["args"]) if cases else None
    methods = signatures(text, cases[0]["args"][0]) if entry else {}
    kinds = {} if entry else guess_kinds(params or [], text, cases[0]["args"] if cases else [], found_kinds)
    for name, kind in kinds.items():  # Kinds the page writes (arrows) need no word; kinds read from a name do.
        if name not in found_kinds and kind in KIND_WORDS and [name] != renamed:
            notes.append(f"{name} is read as {KIND_WORDS[kind]}, because the page calls it {name} in a problem about "
                         f"{'trees' if kind == 'tree' else 'linked lists'}. That is an inference; change its kind if it's wrong.")
    if entry:
        notes.append(f"A design problem: write a class {entry} with the methods it calls ({', '.join(dict.fromkeys(cases[0]['args'][0][1:]))}). Each case runs the operations in order.")
    answer = None
    if any(r.get("node_value") for r in records):
        answer = "node-value"
        notes.append("The page shows the value at the node your function returns, so the lab compares that node's value (None when you return None).")
    elif pattern or (cases and re.search(r"\bprint", text, re.I) and all(isinstance(c["expected"], list) and c["expected"] and all(isinstance(x, str) for x in c["expected"]) for c in cases)):
        answer = "lines"
    if "cycle" in kinds.values():
        notes.append("pos says where the tail links back to (-1 for no cycle): it builds the cycle and isn't passed to your function, as on the judges.")
    return dict(title=title, statement=text, description="\n".join(statement), sections=sections, constraints=constraints, examples=records,
                params=params or [], cases=[{**case, "source": "page"} for case in cases], notes=notes, source=url, kinds=kinds, entry=entry or "solve", methods=methods,
                images=images, truncated=truncated, answer=answer, unfilled=unfilled)


def drawn_lines(raw):
    """A paragraph's lines as a browser draws them, when that is certain: each &nbsp; is a space that stays, while
    ordinary spaces at a line's start or in a run are dropped or merged. None when the page writes such spaces,
    because then what it draws and what it means can differ (a pattern indented with plain spaces)."""
    lines = raw.split("\n")
    while lines and not lines[0].strip():
        lines.pop(0)
    while lines and not lines[-1].strip():
        lines.pop()
    drawn = []
    for line in lines:
        line = line.rstrip()  # Trailing spaces are never compared.
        if re.match(r"[ \t\r\f]", line) or re.search(r"[ \t\r\f]{2}|[\t\r\f]", line):
            return None
        drawn.append(line.replace("\xa0", " "))
    return drawn


def statement_pattern(statement, params, cases, sources=None):
    """The pattern a statement draws in text for one size ("Let's say for N = 5, the pattern should look like as
    below:" followed by its lines), as a case: that size in, the lines out. Only for a problem with one numeric input.
    → (case or None, a note when the pattern is there but its spacing can't be read exactly, {statement index: the
    block as drawn}). In a code block the spacing is exact; elsewhere it is exact only when written with &nbsp;."""
    if not params or len(params) != 1 or (cases and type(cases[0]["args"][0]) is not int):
        return None, None, {}
    looks = lambda parts: all(part.strip() and len(part) <= 60 and not re.search(r"[a-z]{3,}", part) for part in parts)
    for i, line in enumerate(statement):
        # The size and its phrase on the line that ends in a colon (searched there alone, so a line repeating the
        # phrase is read once rather than once per repeat).
        tail = line.rstrip()
        last = tail.rsplit("\n", 1)[-1]
        size = re.search(r"\bn\s*=\s*(\d{1,2})\b[^\n]{0,80}?(?:look like|looks like|as below|as follows|following)", last, re.I) if tail.endswith((":", "：")) else None
        if not size:
            continue
        rows, drawn = [], {}
        for at in range(i + 1, len(statement)):
            if not looks(statement[at].split("\n")):
                break
            source = sources[at] if sources else None
            parts = statement[at].split("\n") if source is None else drawn_lines(source)
            if parts is None:
                return None, (f"The statement draws the pattern for {params[0]} = {size.group(1)} outside a code block with plain spaces, which a browser "
                              "merges or drops, so its exact spacing can't be read and it wasn't made a case. Compare with the original page."), {}
            if not looks(parts):
                break
            rows += [part.rstrip() for part in parts]
            if source is not None:
                drawn[at] = "\n".join(parts)
        if rows:
            return (dict(name="From the statement", args=[int(size.group(1))], expected=rows, missing=False,
                         explanation=f"The pattern the statement draws for {params[0]} = {size.group(1)}."), None, drawn)
    return None, None, {}


MD_LINK = re.compile(r"\[!\[[^\]]*\]\([^)]*\)\]\([^)]*\)|!?\[[^\]]*\]\([^)]*\)")  # [text](url), ![alt](src), [![badge](src)](url)


def markdown_page(text):
    """A problem written in Markdown (a README) as HTML the same reader understands. Inline HTML is kept;
    front matter and lines that are only links (a language switch, badges) are navigation, not the problem."""
    out, fence, buffer, paragraph, raw = [], None, [], [], False

    def flush():
        if paragraph:
            out.append("<p>" + "<br>".join(paragraph) + "</p>")
            paragraph.clear()

    text = re.sub(r"\A---\n.*?\n---\n", "", text.replace("\r\n", "\n"), count=1, flags=re.S)
    # The link patterns scan from every "[" to its "]": brackets × length per line. A page gets a budget of that
    # work; a line past it (only a crafted page gets there) keeps its link syntax as written.
    budget = 4_000_000
    for line in text.split("\n"):
        if fence is not None:
            if line.strip().startswith("```"):
                out.append("<pre>" + html.escape("\n".join(buffer)) + "</pre>")
                fence, buffer = None, []
            else:
                buffer.append(line)
            continue
        if raw:  # Inside an HTML <pre>: every line as written.
            out.append(line)
            raw = "</pre>" not in line.lower()
            continue
        cost = (line.count("[") + 1) * len(line) if "](" in line else 0  # Without "](" no link pattern can match.
        links = 0 < cost <= budget
        budget -= cost if links else 0
        if links and MD_LINK.search(line) and not MD_LINK.sub("", line).strip(" \t|·•-"):
            flush()
            out += [f'<img alt="{html.escape(alt)}">' for alt in re.findall(r"(?<!\[)!\[([^\]]*)\]\(", line)]  # A figure still counts; a badge doesn't.
            continue
        if re.match(r"\s*<pre\b", line, re.I):
            flush()
            out.append(line)
            raw = "</pre>" not in line.lower()
            continue
        if line.strip().startswith("```"):
            flush()
            fence = line.strip()
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        item = re.match(r"^\s*(?:[-*+]|\d+\.)\s+(.*)$", line)
        if heading:
            flush()
            label = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", heading.group(2)) if links else heading.group(2)
            out.append(f"<h{len(heading.group(1))}>{label}</h{len(heading.group(1))}>")
        elif item and not line.lstrip().startswith("<"):
            flush()
            out.append(f"<li>{item.group(1)}</li>")
        elif not line.strip():
            flush()
        elif re.match(r"\s*<(?:/?(?:p|div|ul|ol|li|table|thead|tbody|tr|td|th|h[1-6]|blockquote|img|br|hr|details|summary|section|dl|dt|dd)\b|!--)", line, re.I):
            flush()
            out.append(line)  # Block HTML as written; a line opening with inline HTML (<strong>Follow-up:</strong> …) is paragraph text.
        else:
            line = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", line) if links else line  # A Markdown link reads as its text.
            paragraph.append(re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", re.sub(r"`([^`]+)`", r"<code>\1</code>", line)))
    flush()
    if fence is not None:
        out.append("<pre>" + html.escape("\n".join(buffer)) + "</pre>")
    return "\n".join(out)


ROBOTS = {}


def allowed(url, fetcher):
    """Respect the site's robots.txt for this page (cached per site), as Python's own robot parser reads the answer:
    no robots.txt (404 and other 4xx) allows everything, 401/403 disallow everything. When it can't be read at all
    (a server error, no connection) the page can't be confirmed as allowed, so it isn't read, and nothing is cached."""
    parts = urllib.parse.urlsplit(url)
    root = f"{parts.scheme}://{parts.netloc}"
    rules = ROBOTS.get(root)
    if rules is None:
        rules = urllib.robotparser.RobotFileParser()
        token = REDIRECT_CHECK.set(None)  # robots.txt itself is read without the page's redirect check.
        try:
            body, _, _, charset = fetcher(root + "/robots.txt")
        except ValueError as error:
            status = getattr(error, "status", None)
            if status in (401, 403):
                rules.disallow_all = True
            elif status is not None and 400 <= status < 500:
                rules.allow_all = True
            else:
                reason = f"the site answered {status}" if status else "the site couldn't be reached"
                raise ValueError(f"its robots.txt couldn't be read ({reason}), so it can't be confirmed that the page may be read")
        else:
            rules.parse(decode(body, charset).splitlines())
        finally:
            REDIRECT_CHECK.reset(token)
        ROBOTS[root] = rules
        while len(ROBOTS) > 64:
            ROBOTS.pop(next(iter(ROBOTS)))
    return rules.can_fetch(AGENT, url)


def is_leetcode(url):
    host = urllib.parse.urlsplit(url).hostname or ""
    return host == "leetcode.com" or host.endswith(".leetcode.com")


def read_problem(row, fetcher=None):
    """Read one sheet row's problem, on request. The sheet's own site comes first (its page for the
    problem); only if the problem isn't readable there are the row's attached links tried, in order.
    leetcode.com itself is never fetched (its pages are built in the browser behind a bot check); when the row
    links a LeetCode problem, LeetCode's own version of it (leetcode.py) completes what the page leaves out, and is
    the problem when no page can be read. Nothing is guessed: every value says where it came from."""
    fetcher = fetcher or fetch
    candidates = []
    for link in [row.get("source"), row.get("url"), *(row.get("links") or [])]:
        if isinstance(link, str) and link.strip() and link.strip() not in candidates:
            candidates.append(link.strip())
    slug = leetcode.row_slug({"source": row.get("source"), "url": row.get("url"), "links": row.get("links")})
    if not candidates:
        raise ValueError("This problem has no link to read it from. Paste its statement and examples instead.")
    misses = []
    for target in candidates:
        host = (urllib.parse.urlsplit(target).hostname or "").removeprefix("www.")
        if is_leetcode(target):
            misses.append(f"{host} builds its pages in the browser, so it can't be read here")
            continue
        try:
            if not allowed(target, fetcher):
                raise ValueError(f"{host} asks automated tools not to read that page (robots.txt)")
            token = REDIRECT_CHECK.set(lambda url: allowed(url, fetcher))  # A redirect must be allowed too.
            try:
                body, final, kind, charset = fetcher(target)
            finally:
                REDIRECT_CHECK.reset(token)
            if final != target and (is_leetcode(final) or not allowed(final, fetcher)):  # However the fetcher followed it.
                raise ValueError(f"the page redirected to {urllib.parse.urlsplit(final).hostname}, which can't be read here")
            text = decode(body, charset)
            if kind in ("text/markdown", "text/x-markdown") or urllib.parse.urlsplit(final).path.lower().endswith((".md", ".markdown")):
                problem = problem_from_page(markdown_page(text), final)  # A problem written as a README.
            elif kind in ("text/html", "application/xhtml+xml"):
                problem = problem_from_page(text, final)
            else:
                raise ValueError("that link isn't a problem page")
        except ValueError as error:
            misses.append(f"{host}: {str(error).rstrip('.')}")
            continue
        if unknown_charset(charset):
            problem["notes"].insert(0, f"The page names a character set that can't be read ({charset[:40]}), so it was read as UTF-8. Check any unusual characters against the original.")
        if misses:
            first = (urllib.parse.urlsplit(candidates[0]).hostname or "").removeprefix("www.")
            problem["notes"].insert(0, f"The problem couldn't be read on {first} ({misses[0]}), so it was read from its attached link on {host}.")
        return complete_from_leetcode(problem, slug, fetcher)
    if slug:  # No page could be read: LeetCode's own version of the problem, if the row links one.
        try:
            found = leetcode.page(slug, fetcher, decode)
        except ValueError as error:
            misses.append(f"LeetCode's version couldn't be read ({str(error).rstrip('.')})")
        else:
            if found:
                read = leetcode.problem(found, problem_from_page)
                read.pop("unfilled", None)
                return read
            misses.append("LeetCode's version of this problem isn't in its open mirror")
    raise ValueError("The problem couldn't be read: " + "; ".join(misses) + ". Open it and paste its statement and examples instead.")


def complete_from_leetcode(problem, slug, fetcher):
    """The page's problem, completed from LeetCode's version where the page leaves something out and the two
    provably describe the same problem:
    - an example whose output the page doesn't give, when LeetCode states an example with exactly the same input;
    - the parameters' kinds, and an in-place answer, from LeetCode's signature when it names the same parameters;
    - every example, when the page has none (marked as LeetCode's).
    Where both give an output for the same input they must agree, or nothing of LeetCode's is used. The page's
    own values are never replaced."""
    unfilled = problem.pop("unfilled", [])
    if not slug:
        return problem
    try:
        found = leetcode.page(slug, fetcher, decode)
    except ValueError as error:
        problem["notes"].append(f"LeetCode's version of this problem couldn't be read to complete it ({str(error).rstrip('.')}).")
        return problem
    if not found:
        return problem
    theirs = leetcode.problem(found, problem_from_page)
    stated = {json.dumps(case["args"], sort_keys=True): case for case in theirs["cases"] if not case.get("missing")}
    ours = [case for case in problem["cases"] if not case.get("missing")]
    clash = next((case for case in ours if json.dumps(case["args"], sort_keys=True) in stated and not same_output(stated[json.dumps(case["args"], sort_keys=True)]["expected"], case["expected"])), None)
    if clash:
        problem["notes"].append(f"LeetCode's version gives a different output for {clash['name']}'s input, so it isn't the same contract and nothing of it was used.")
        return problem
    problem["leetcode"] = slug
    filled = []
    for case in [c for c in problem["cases"] if c.get("missing")] + unfilled:
        match = stated.get(json.dumps(case["args"], sort_keys=True))
        if match:
            case.update(expected=match["expected"], missing=False, source="leetcode")
            filled.append(case)
            if case in unfilled:
                problem["cases"].append(case)
    if filled:
        problem["notes"].append(f"The output of {', '.join(c['name'] for c in filled)} is LeetCode's, for exactly the same input ({leetcode.CREDIT}): the page doesn't give it as a value.")
    if not problem["cases"] and theirs["cases"] and (not problem["params"] or len(problem["params"]) == len(theirs["params"])):
        problem["cases"] = theirs["cases"]
        problem["params"] = problem["params"] or theirs["params"]
        problem["notes"].append(f"The page has no examples to read, so these are LeetCode's examples of the same problem ({leetcode.CREDIT}). Check its contract matches the page's.")
    if leetcode.apply_signature(problem, found):
        problem["notes"].append(f"Parameter kinds come from LeetCode's Python signature for this problem: {', '.join(f'{n}: {t}' for n, t in found['signature']['params'])}" + (" (it changes its input in place and returns nothing)." if found["signature"]["returns"] == "None" else "."))
    return problem


def same_output(a, b):
    """Two stated outputs agree: equal values, numbers within the judges' tolerance."""
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) <= 1e-5 * max(1, abs(a), abs(b))
    return json.dumps(a, sort_keys=True) == json.dumps(b, sort_keys=True)


def describe(value):
    return {bool: "true/false", int: "an integer", float: "a number", str: "text", list: "a list", dict: "a dictionary"}.get(type(value), "None")


def build_lab(data):
    """Validate a learner-defined lab. Its cases are the learner's own: the expected outputs are
    their specification, and the lab says so wherever it judges against them."""
    if not isinstance(data, dict):
        raise ValueError("Send the lab as an object.")
    title = clean_title(data.get("title") if isinstance(data.get("title"), str) else "")
    statement = data.get("statement") if isinstance(data.get("statement"), str) else ""
    returns = data.get("returns") if isinstance(data.get("returns"), str) else ""
    if len(title) < 2:
        raise ValueError("Give the lab a title.")
    if len(statement.strip()) < 10 or len(statement) > MAX_STATEMENT:
        raise ValueError(f"Describe the problem in 10 to {MAX_STATEMENT:,} characters.")
    if len(returns) > 400:
        raise ValueError("Keep the return description under 400 characters.")
    params = data.get("params")
    if isinstance(params, str):
        params = [p.strip() for p in params.split(",") if p.strip()]
    if not isinstance(params, list) or not 1 <= len(params) <= 5 or any(not isinstance(p, str) or not re.fullmatch(r"[A-Za-z_]\w{0,24}", p) for p in params):
        raise ValueError("Name 1 to 5 parameters, separated by commas, e.g. nums, target.")
    if len(set(params)) != len(params) or any(p in {"solve", "self"} or keyword.iskeyword(p) for p in params):
        raise ValueError("Use distinct parameter names that aren't Python keywords.")
    kinds = data.get("kinds") if isinstance(data.get("kinds"), dict) else {}
    kinds = {k: v for k, v in kinds.items() if k in params and v in KINDS}
    if "cycle" in kinds.values() and not any(v in ("linkedlist", "dll") for v in kinds.values()):
        raise ValueError("A cycle position needs a linked list to link back into.")
    answer = data.get("answer") if data.get("answer") in ANSWERS else None
    entry = data.get("entry") if isinstance(data.get("entry"), str) and data.get("entry") else "solve"
    if entry != "solve" and (not re.fullmatch(r"[A-Z]\w{0,40}", entry) or len(params) != 2):
        raise ValueError("A design problem has a class name and two inputs: the operations and their arguments.")
    cases = data.get("cases")
    if not isinstance(cases, list) or not 1 <= len(cases) <= MAX_CASES:
        raise ValueError(f"Add 1 to {MAX_CASES} cases.")
    clean = []
    for number, case in enumerate(cases, 1):
        if not isinstance(case, dict) or not isinstance(case.get("args"), list) or len(case["args"]) != len(params):
            raise ValueError(f"Case {number} needs one value for each parameter ({', '.join(params)}).")
        name = clean_title(case.get("name") if isinstance(case.get("name"), str) else "") or f"Case {number}"
        try:
            args = []
            for param, value in zip(params, case["args"]):
                try:
                    args.append(literal(value) if isinstance(value, str) else plain(value))
                except ValueError as error:
                    raise ValueError(f"{param}: {error}")
            if "expected" not in case or case["expected"] is None or (isinstance(case["expected"], str) and not case["expected"].strip()):
                raise ValueError("it needs an expected output, which is how the lab knows when your code is on target.")
            try:
                expected = literal(case["expected"]) if isinstance(case["expected"], str) else plain(case["expected"])
            except ValueError as error:
                raise ValueError(f"expected output: {error}")
        except ValueError as error:
            raise ValueError(f"{name} — {error}")
        explanation = case.get("explanation") if isinstance(case.get("explanation"), str) else ""
        source = case.get("source") if case.get("source") in SOURCES else None
        clean.append(dict(name=name[:40], args=args, expected=expected, **({"explanation": explanation.strip()[:1000]} if explanation.strip() else {}), **({"source": source} if source else {})))
    shapes = [type(v) for v in clean[0]["args"]]
    for number, case in enumerate(clean[1:], 2):
        for name, value, shape in zip(params, case["args"], shapes):
            if type(value) is not shape and not ({type(value), shape} <= {int, float}):
                raise ValueError(f"Case {number}: {name} is {describe(value)}, but case 1 gives {describe(clean[0]['args'][params.index(name)])}. Keep each parameter's type.")
    names = [c["name"] for c in clean]
    for i, name in enumerate(names):
        if names.count(name) > 1:
            clean[i]["name"] = f"{name} ({i + 1})"
    if entry != "solve":
        for number, case in enumerate(clean, 1):
            operations, arguments = case["args"]
            if not isinstance(operations, list) or not operations or operations[0] != entry or any(not isinstance(o, str) or not re.fullmatch(r"[A-Za-z]\w*", o) for o in operations):
                raise ValueError(f"Case {number}: the operations must be method names, starting with {entry}.")
            if not isinstance(arguments, list) or len(arguments) != len(operations):
                raise ValueError(f"Case {number}: give one argument entry for each operation.")
    order = "any" if data.get("order") == "any" else "exact"
    difficulty = difficulty_of(data.get("difficulty")) or ""
    topic = clean_title(data.get("topic") if isinstance(data.get("topic"), str) else "")[:80]
    methods = data.get("methods") if isinstance(data.get("methods"), dict) else {}
    methods = {k: [a for a in v if isinstance(a, str) and re.fullmatch(r"[A-Za-z_]\w*", a)][:6] for k, v in methods.items() if isinstance(k, str) and isinstance(v, list)}
    source = data.get("source") if isinstance(data.get("source"), str) and re.match(r"https?://", data.get("source") or "") else ""
    slug = data.get("leetcode") if isinstance(data.get("leetcode"), str) and re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+){0,15}", data.get("leetcode") or "") else None
    return dict(title=title, statement=statement.strip(), returns=returns.strip(), params=params, cases=clean, order=order, difficulty=difficulty, topic=topic,
                kinds=kinds, entry=entry, methods=methods if entry != "solve" else {}, source=source[:500], **({"answer": answer} if answer else {}), **({"leetcode": slug} if slug else {}))


KIND_NOTES = {"linkedlist": "a linked list: a ListNode with .val and .next (None at the end)",
              "dll": "a doubly linked list: a ListNode with .val, .next and .prev",
              "tree": "a binary tree: a TreeNode with .val, .left and .right (None for no child)",
              "linkedlists": "a list of linked lists: each item is the head ListNode of one list"}
ANSWER_NOTES = {"in-place": "    # Change the first input in place: the lab checks it as solve leaves it (what solve returns is ignored).",
                "node-value": "    # Return the node itself (or None): the lab compares its value, as the page shows it.",
                "lines": "    # Return the printed lines as a list of strings; trailing spaces don't matter."}


def starter_code(lab):
    """A starter that names what each input really is; for a design problem, the class to write."""
    entry = lab.get("entry", "solve")
    if entry != "solve":
        operations, arguments = lab["cases"][0]["args"]
        methods, lines = lab.get("methods") or {}, [f"class {entry}:", f"    # {lab['title']}: each case calls these methods in order."]
        for op in dict.fromkeys(operations):
            index = operations.index(op)
            given = arguments[index] if index < len(arguments) else []
            count = len(given) if isinstance(given, list) else 1
            names = methods.get(op) or ([f"arg{i + 1}" for i in range(count)] if count > 1 else ["value"] if count else [])
            name = "__init__" if op == entry else op
            lines += ["", f"    def {name}({', '.join(['self'] + names)}):", "        pass"]
        return "\n".join(lines) + "\n"
    kinds = lab.get("kinds", {})
    notes = [f"    # {name} is {KIND_NOTES[kind]}." for name, kind in kinds.items() if kind in KIND_NOTES]
    returns = [ANSWER_NOTES[lab["answer"]]] if lab.get("answer") in ANSWER_NOTES else ["    # Return the head (or root) of your answer: it is compared as a list of values."] if any(k in ("linkedlist", "dll", "tree") for k in kinds.values()) else []
    # A cycle position only builds the input; the function receives the list it describes.
    params = [name for name in lab["params"] if kinds.get(name) != "cycle"]
    return "\n".join([f"def solve({', '.join(params)}):", f"    # {lab['title']}", *notes, *returns, "    # Write your approach here.", "    ", "    pass"]) + "\n"


RECALL = ["Without looking at your code: what does your solution keep track of as it moves through the input?",
          "Which of your cases was the hardest to get right, and what did it teach you about the problem?",
          "Name an input none of your cases covers. What should your solution return for it, and why?"]


def as_problem(lab_id, lab, sheet_id, sheet_name, number):
    """A learner-defined lab in the shape every part of the laboratory reads."""
    params, first = lab["params"], lab["cases"][0]
    given = "; ".join(f"{name}: {describe(value)}" for name, value in zip(params, first["args"]))
    returns = lab["returns"] or f"Return {describe(first['expected'])}, compared with the expected output of each of your cases."
    if lab["order"] == "any":
        returns += " The order of the returned items doesn't matter."
    starter = starter_code(lab)
    return dict(id=lab_id, custom=True, sheetId=sheet_id, sheetName=sheet_name, number=number, title=lab["title"],
                category=lab["topic"] or "Your sheet", difficulty=lab["difficulty"] or "Your lab", params=params,
                statement=lab["statement"], decoder=dict(given=given, find=lab["statement"][:400], returns=returns),
                example=dict(args=first["args"], expected=first["expected"]), tests=lab["cases"], cases=lab["cases"], order=lab["order"],
                returnsNote=lab["returns"], starter=starter, discovery=[], hints=[], recall=RECALL, transfer=None, complexity=None,
                kinds=lab.get("kinds", {}), entry=lab.get("entry", "solve"), methods=lab.get("methods", {}), sourceUrl=lab.get("source", ""), answer=lab.get("answer"), leetcode=lab.get("leetcode"),
                solution=None, brute=None)
